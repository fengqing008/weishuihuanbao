#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""国标公文智能处理助手 - 格式化导出工具 v18.5
解析 Markdown → 生成符合 GB/T 9704-2012 的 Word 文档。
v18.5：匿名化通用版 | v18.3：主送机关规范 | v18.2：性能优化版 | v18.1：附件表格编排 | v18：正文表格排版"""
import argparse, os, re, sys, zipfile, shutil
from copy import deepcopy
from datetime import datetime
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from lxml import etree

# ============ 路径常量 ============
_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_SKILL_DIR = os.path.dirname(_SCRIPT_DIR)
_ASSETS_DIR = os.path.join(_SKILL_DIR, "assets")

def _get_desktop_path():
    home = os.path.expanduser("~")
    for name in ("Desktop", "\u684c\u9762"):
        p = os.path.join(home, name)
        if os.path.isdir(p): return p
    return home

_DEFAULT_OUTPUT = _get_desktop_path()
_FONT_FANGSONG = os.path.join(_ASSETS_DIR, "仿宋_GB2312.ttf")
_FONT_BIAOSONG = os.path.join(_ASSETS_DIR, "方正小标宋简体.ttf")

# ============ 预编译常量 ============
FONT_BLACK, FONT_KAI, FONT_SONG = "黑体", "楷体_GB2312", "仿宋_GB2312"
FONT_BIAOSONG_NAME, FONT_ENGLISH, FONT_SONG_TI = "方正小标宋简体", "Times New Roman", "宋体"
PT_36, PT_18, PT_16, PT_14, PT_28, PT_0 = Pt(36), Pt(18), Pt(16), Pt(14), Pt(28), Pt(0)
INDENT_2CHAR = Pt(32)

# 预编译正则
_RE_TITLE_SENDER = re.compile(r'^(.+?(?:学院|大学|处|部|中心|办公室))(关于)')
_RE_DATE = re.compile(r'^\d{4}年\d{1,2}月\d{1,2}日$')
_RE_ATTACH_NUM = re.compile(r'^\d+\.')
_RE_TABLE_CAPTION = re.compile(r'^表\s*\d+')
_RE_TABLE_ROW = re.compile(r'^\s*\|.+\|\s*$')
_RE_TABLE_SEP = re.compile(r'^\s*\|[\s\-:]+\|')
_RE_IS_NUMERIC = re.compile(r'^[\d,，.%\s\-]+$')
_ENGLISH_RE = re.compile(r'([a-zA-Z]+)')

# 标点加速表
_PUNCT_TABLE = str.maketrans({',':'\uff0c',';':'\uff1b','(':'\uff08',')':'\uff09',':':'\uff1a','?':'\uff1f','!':'\uff01'})

# Markdown加粗剥离
_MD_BOLD_ASTERISK = re.compile(r'\*\*(.+?)\*\*')
_MD_BOLD_UNDERSCORE = re.compile(r'__(.+?)__')

def _strip_markdown(text):
    return _MD_BOLD_UNDERSCORE.sub(r'\1', _MD_BOLD_ASTERISK.sub(r'\1', text))

def _normalize_punctuation(text):
    text = text.translate(_PUNCT_TABLE)
    result, dq, sq = [], True, True
    for ch in text:
        if ch == '"': result.append('\u201c' if dq else '\u201d'); dq = not dq
        elif ch == "'": result.append('\u2018' if sq else '\u2019'); sq = not sq
        else: result.append(ch)
    return ''.join(result)

# ============ XML 元素缓存 ============
def _make_rfonts(font_name):
    e = OxmlElement('w:rFonts')
    for a in ('eastAsia','ascii','hAnsi'): e.set(qn(f'w:{a}'), font_name)
    return e

_RFONT_CACHE = {}
def _get_rfont(font_name):
    if font_name not in _RFONT_CACHE: _RFONT_CACHE[font_name] = _make_rfonts(font_name)
    return deepcopy(_RFONT_CACHE[font_name])

_XML_CACHE = {}
def _cached_xml(tag, val):
    k = (tag, val)
    if k not in _XML_CACHE:
        e = OxmlElement(tag); e.set(qn('w:val'), val); _XML_CACHE[k] = e
    return deepcopy(_XML_CACHE[k])

def _apply_font(run, font_name, font_size, bold=False):
    rPr = run._r.find(qn('w:rPr'))
    if rPr is None: rPr = OxmlElement('w:rPr'); run._r.insert(0, rPr)
    for tag in ('w:rFonts','w:sz','w:b','w:szCs','w:bCs'):
        for old in rPr.findall(qn(tag)): rPr.remove(old)
    rPr.append(_get_rfont(font_name))
    hv = str(int(font_size.pt * 2)); bv = '1' if bold else '0'
    rPr.append(_cached_xml('w:sz', hv)); rPr.append(_cached_xml('w:szCs', hv))
    rPr.append(_cached_xml('w:b', bv)); rPr.append(_cached_xml('w:bCs', bv))

def _set_para_fmt(para, alignment, line_spacing=PT_28, space_before=PT_0,
                  space_after=PT_0, first_indent=None):
    para.alignment = alignment
    pf = para.paragraph_format
    pf.line_spacing = line_spacing; pf.space_before = space_before
    pf.space_after = space_after
    if first_indent is not None:
        if first_indent is not None:
            if int(first_indent) == 0:
                # Pt(0) → 显式设置 firstLine=0，确保主送机关顶格（覆盖 Normal 样式默认缩进）
                pPr = para._element.get_or_add_pPr()
                existing_ind = pPr.find(qn('w:ind'))
                if existing_ind is not None:
                    pPr.remove(existing_ind)
                ind = OxmlElement('w:ind')
                ind.set(qn('w:firstLine'), '0')
                pPr.append(ind)
            else:
                pf.first_line_indent = first_indent

def _add_para(doc, text, font_name, font_size, alignment=WD_ALIGN_PARAGRAPH.LEFT,
              first_indent=None, bold=False, space_before=PT_0, space_after=PT_0):
    text = _normalize_punctuation(text)
    p = doc.add_paragraph()
    _set_para_fmt(p, alignment, space_before=space_before, space_after=space_after,
                  first_indent=first_indent)
    for seg in _ENGLISH_RE.split(text):
        if not seg: continue
        is_en = bool(_ENGLISH_RE.fullmatch(seg))
        r = p.add_run(seg); _apply_font(r, FONT_ENGLISH if is_en else font_name, font_size, bold)
    return p

# ============ 公文标题样式定义（v18.6：标题挂大纲级别，可一键生成目录） ============
def _set_style_font(style, cn_font, size, bold=False, align=None):
    """配置某个 Word 样式的公文字体/字号/颜色；清除主题字体等非公文属性"""
    style.font.size = size
    style.font.bold = bold
    try:
        style.font.color.rgb = RGBColor(0, 0, 0)
    except Exception:
        pass
    rPr = style.element.get_or_add_rPr()
    rf = rPr.get_or_add_rFonts()
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    rf.set(qn('w:ascii'), FONT_ENGLISH)
    rf.set(qn('w:hAnsi'), FONT_ENGLISH)
    rf.set(qn('w:cs'), FONT_ENGLISH)
    rf.set(qn('w:eastAsia'), cn_font)
    pf = style.paragraph_format
    pf.space_before = PT_0
    pf.space_after = PT_0
    pf.line_spacing = PT_28
    pf.first_line_indent = Pt(0)
    if align is not None:
        pf.alignment = align


def setup_gw_styles(doc):
    """定义 GB/T 9704-2012 公文标准样式。
    Title    = 方正小标宋简体 二号 居中（文档大标题，不计入目录）
    Heading1 = 黑体 三号        （一级标题「一、」）→ 大纲级别 1
    Heading2 = 楷体_GB2312 三号 （二级标题「（一）」）→ 大纲级别 2
    Heading3 = 仿宋_GB2312 三号 加粗（三级标题「1.」）→ 大纲级别 3
    Normal   = 仿宋_GB2312 三号 （正文）
    说明：Heading 1/2/3 是 Word 内置样式，自带 w:outlineLvl，故用其排版的标题
    可被「引用→目录」一键识别生成目录。"""
    _spec = {
        'Normal':    (FONT_SONG, PT_16, False, None),
        'Title':     (FONT_BIAOSONG_NAME, Pt(22), False, WD_ALIGN_PARAGRAPH.CENTER),
        'Heading 1': (FONT_BLACK, PT_16, False, None),
        'Heading 2': (FONT_KAI, PT_16, False, None),
        'Heading 3': (FONT_SONG, PT_16, True, None),
    }
    for name, (cf, sz, bd, al) in _spec.items():
        try:
            _set_style_font(doc.styles[name], cf, sz, bold=bd, align=al)
        except Exception:
            pass


def _add_heading_para(doc, text, level):
    """按公文标题层级添加段落并挂接 Word 大纲级别。
    level: 'title'→Title 样式；1/2/3→Heading 1/2/3 样式。
    段落挂样式（保证大纲级别）＋ run 级设字体（保证字体严格国标，双保险）。"""
    text = _normalize_punctuation(text)
    style_map = {1: 'Heading 1', 2: 'Heading 2', 3: 'Heading 3', 'title': 'Title'}
    cfg = {1: (FONT_BLACK, PT_16, False), 2: (FONT_KAI, PT_16, False),
           3: (FONT_SONG, PT_16, True), 'title': (FONT_BIAOSONG_NAME, Pt(22), False)}
    sname = style_map.get(level)
    p = None
    if sname:
        try:
            p = doc.add_paragraph(style=sname)
        except Exception:
            p = None
    if p is None:
        p = doc.add_paragraph()
        if sname:
            try:
                p.style = doc.styles[sname]
            except Exception:
                pass
    if level == 'title':
        _set_para_fmt(p, WD_ALIGN_PARAGRAPH.CENTER, first_indent=Pt(0))
    else:
        # 一级/二级/三级标题 左空二字（首行缩进 2 字符，某负责人 2026-09-12 定）
        _set_para_fmt(p, WD_ALIGN_PARAGRAPH.LEFT, first_indent=INDENT_2CHAR)
    fn, fs, bd = cfg.get(level, (FONT_SONG, PT_16, False))
    for seg in _ENGLISH_RE.split(text):
        if not seg:
            continue
        is_en = bool(_ENGLISH_RE.fullmatch(seg))
        r = p.add_run(seg)
        _apply_font(r, FONT_ENGLISH if is_en else fn, fs, bd)
    return p


def _add_empty_line(doc, font_name=FONT_SONG, font_size=PT_16, count=1):
    for _ in range(count):
        p = doc.add_paragraph(); _set_para_fmt(p, WD_ALIGN_PARAGRAPH.LEFT)
        r = p.add_run(""); _apply_font(r, font_name, font_size, False)

# ============ 表格排版 ============
def _is_numeric_cell(text):
    return bool(text) and bool(_RE_IS_NUMERIC.match(text)) and any(ch.isdigit() for ch in text)

def _set_cell_text(cell, text, font_name, font_size, alignment=WD_ALIGN_PARAGRAPH.LEFT,
                   bold=False, line_spacing=PT_28):
    for p in cell.paragraphs:
        for r in p.runs: r.text = ''
    para = cell.paragraphs[0]; para.alignment = alignment
    pf = para.paragraph_format
    pf.line_spacing = line_spacing; pf.space_before = PT_0; pf.space_after = PT_0
    text = _normalize_punctuation(text)
    for seg in _ENGLISH_RE.split(text):
        if not seg: continue
        is_en = bool(_ENGLISH_RE.fullmatch(seg))
        r = para.add_run(seg)
        _apply_font(r, FONT_ENGLISH if is_en else font_name, font_size, bold)

def _apply_table_borders(table):
    """GB/T公文表格边框：顶/底/表头分隔线=1pt粗实线，栏线/行线=0.5pt细实线，表头跨页重复"""
    tbl = table._tbl
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is None: tblPr = OxmlElement('w:tblPr'); tbl.insert(0, tblPr)

    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None: tblW = OxmlElement('w:tblW'); tblPr.append(tblW)
    tblW.set(qn('w:w'), '8849'); tblW.set(qn('w:type'), 'dxa')

    tblJc = tblPr.find(qn('w:jc'))
    if tblJc is None: tblJc = OxmlElement('w:jc'); tblPr.append(tblJc)
    tblJc.set(qn('w:val'), 'center')

    tblBorders = OxmlElement('w:tblBorders')
    for bn, sz in [('top','8'),('bottom','8'),('left','8'),('right','8'),
                   ('insideH','4'),('insideV','4')]:
        b = OxmlElement(f'w:{bn}')
        for a,v in [('val','single'),('sz',sz),('space','0'),('color','000000')]:
            b.set(qn(f'w:{a}'), v)
        tblBorders.append(b)
    tblPr.append(tblBorders)

    header_row = tbl.find(qn('w:tr'))
    if header_row is not None:
        header_trPr = header_row.find(qn('w:trPr'))
        if header_trPr is None: header_trPr = OxmlElement('w:trPr'); header_row.insert(0, header_trPr)
        thdr = OxmlElement('w:tblHeader'); thdr.set(qn('w:val'), 'true')
        header_trPr.append(thdr)
        for tc in header_row.findall(qn('w:tc')):
            tcPr = tc.find(qn('w:tcPr'))
            if tcPr is None: tcPr = OxmlElement('w:tcPr'); tc.insert(0, tcPr)
            tcBrd = tcPr.find(qn('w:tcBorders'))
            if tcBrd is None: tcBrd = OxmlElement('w:tcBorders'); tcPr.append(tcBrd)
            btm = OxmlElement('w:bottom')
            for a,v in [('val','single'),('sz','8'),('space','0'),('color','000000')]:
                btm.set(qn(f'w:{a}'), v)
            tcBrd.append(btm)

def _render_table(doc, table_data):
    headers = table_data.get('headers', [])
    rows = table_data.get('rows', [])
    if not headers or not rows: return
    caption = table_data.get('caption')
    is_attachment = table_data.get('is_attachment', False)
    attachment_label = table_data.get('attachment_label', '')

    if is_attachment:
        p_page = doc.add_paragraph()
        pf = p_page.paragraph_format
        pf.space_before = PT_0; pf.space_after = PT_0; pf.line_spacing = PT_28
        br_el = OxmlElement('w:br'); br_el.set(qn('w:type'), 'page')
        p_page.add_run('')._element.append(br_el)
        if attachment_label:
            _add_para(doc, attachment_label, FONT_BLACK, PT_16, alignment=WD_ALIGN_PARAGRAPH.LEFT)

    _add_empty_line(doc)
    if caption:
        _add_para(doc, caption, FONT_BLACK, PT_16, alignment=WD_ALIGN_PARAGRAPH.CENTER)

    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ci, h in enumerate(headers):
        _set_cell_text(table.rows[0].cells[ci], h, FONT_BLACK, PT_16, WD_ALIGN_PARAGRAPH.CENTER)
    for ri, row in enumerate(rows):
        for ci, ct in enumerate(row):
            al = WD_ALIGN_PARAGRAPH.RIGHT if _is_numeric_cell(ct) else WD_ALIGN_PARAGRAPH.LEFT
            _set_cell_text(table.rows[ri+1].cells[ci], ct, FONT_SONG, PT_16, al)
    _apply_table_borders(table)
    _add_empty_line(doc)

# ============ 页码：VML 浮动文本框 + outside 定位 ============
_VML_NS, _OVML_NS = 'urn:schemas-microsoft-com:vml', 'urn:schemas-microsoft-com:office:office'
_FOOTER_TEMPLATE = None

def _make_page_number_footer():
    """VML浮动文本框 outside定位：奇数页右侧=右下角，偶数页左侧=左下角"""
    page_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rPr.append(_get_rfont(FONT_SONG_TI))
    rPr.append(_cached_xml('w:sz', str(int(PT_14.pt*2))))
    rPr.append(_cached_xml('w:szCs', str(int(PT_14.pt*2))))
    page_run.append(rPr)
    for ct, at in [('begin',None),('instrText',' PAGE '),('separate',None),('t','1'),('end',None)]:
        if ct in ('begin','separate','end'):
            e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), ct)
        elif ct == 'instrText':
            e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = at
        else:
            e = OxmlElement('w:t'); e.text = at
        page_run.append(e)

    txbx_para = OxmlElement('w:p'); txbx_para.append(OxmlElement('w:pPr')); txbx_para.append(page_run)
    txbx_content = OxmlElement('w:txbxContent'); txbx_content.append(txbx_para)

    shape = etree.Element('{%s}shape'%_VML_NS)
    shape.set('id','_x0000_i1026'); shape.set('type','#_x0000_t202')
    shape.set('style','position:absolute;left:0pt;margin-top:0pt;height:14.25pt;width:36pt;'
              'mso-position-horizontal:outside;mso-position-horizontal-relative:margin;'
              'mso-wrap-style:none;z-index:251659264')
    shape.set('filled','f'); shape.set('stroked','f'); shape.set('coordsize','21600,21600')
    for tag, attrs in [('fill',{'on':'f','focussize':'0,0'}),('stroke',{'on':'f'})]:
        e = etree.SubElement(shape, '{%s}%s'%(_VML_NS,tag))
        for k,v in attrs.items(): e.set(k,v)
    lock = etree.SubElement(shape, '{%s}lock'%_OVML_NS)
    lock.set('{%s}ext'%_VML_NS,'edit'); lock.set('aspectratio','f')
    textbox = etree.SubElement(shape, '{%s}textbox'%_VML_NS)
    textbox.set('inset','0mm,0mm,0mm,0mm'); textbox.set('style','mso-fit-shape-to-text:t')
    textbox.append(txbx_content)

    pict = OxmlElement('w:pict'); pict.append(shape)
    outer = OxmlElement('w:p'); outer.append(OxmlElement('w:pPr'))
    outer_r = OxmlElement('w:r'); outer_r.append(pict); outer.append(outer_r)
    return outer

def _get_footer():
    global _FOOTER_TEMPLATE
    if _FOOTER_TEMPLATE is None: _FOOTER_TEMPLATE = _make_page_number_footer()
    return deepcopy(_FOOTER_TEMPLATE)

def _add_page_number(doc):
    for section in doc.sections:
        sp = section._sectPr
        pnt = sp.find(qn('w:pgNumType'))
        if pnt is None:
            pnt = OxmlElement('w:pgNumType'); cols = sp.find(qn('w:cols'))
            if cols is not None: sp.insert(list(sp).index(cols), pnt)
            else: sp.append(pnt)
        pnt.set(qn('w:fmt'),'numberInDash'); pnt.set(qn('w:start'),'1')
        pgm = sp.find(qn('w:pgMar'))
        if pgm is not None: pgm.set(qn('w:footer'),'992')
        footer = section.footer; footer.is_linked_to_previous = False
        for ep in footer._element.findall(qn('w:p')): footer._element.remove(ep)
        footer._element.append(_get_footer())

def _setup_page(doc):
    for s in doc.sections:
        s.top_margin=Cm(3.7); s.bottom_margin=Cm(3.5); s.left_margin=Cm(2.8)
        s.right_margin=Cm(2.6); s.page_width=Cm(21.0); s.page_height=Cm(29.7)
    setup_gw_styles(doc)
def _embed_fonts(output_path):
    """嵌入字体（调用 gw_font_embed：fontTable + fontTable.rels + settings 四处齐全）"""
    try:
        import sys as _sys
        _here = os.path.dirname(os.path.abspath(__file__))
        if _here not in _sys.path:
            _sys.path.insert(0, _here)
        from gw_font_embed import embed_fonts as _embed
        _fmap = {}
        if os.path.isfile(_FONT_FANGSONG):
            _fmap['仿宋_GB2312'] = _FONT_FANGSONG
        if os.path.isfile(_FONT_BIAOSONG):
            _fmap['方正小标宋简体'] = _FONT_BIAOSONG
        if _fmap:
            _embed(output_path, _fmap)
    except Exception:
        pass
# ============ Markdown 解析 ============
def _parse_table_block(lines, start_idx):
    idx = start_idx; caption = None
    line = lines[idx].strip()
    if _RE_TABLE_CAPTION.match(line): caption = line; idx += 1
    if idx + 2 >= len(lines): return None, start_idx
    hl = lines[idx].strip(); sl = lines[idx+1].strip()
    if not _RE_TABLE_ROW.match(hl) or not _RE_TABLE_SEP.match(sl): return None, start_idx
    headers = [c.strip() for c in hl.strip('| ').split('|')]
    if not headers or all(h=='' for h in headers): return None, start_idx
    rows = []; idx += 2
    while idx < len(lines):
        line = lines[idx].strip()
        if not line: idx += 1; break
        if not _RE_TABLE_ROW.match(line): break
        rows.append([c.strip() for c in line.strip('| ').split('|')])
        idx += 1
    if not rows: return None, start_idx
    return {"type":"table","caption":caption,"headers":headers,"rows":rows}, idx

def parse_markdown(content, doc_type=None, doc_number=None, sender=None,
                   recipients=None, contact=None):
    result = {'header':None,'doc_number':doc_number,'title':None,
              'recipients':recipients,'body_lines':[],'attachments':[],
              'signature':None,'contact':contact,'attachment_tables':[]}
    lines = content.strip().split('\n')
    in_sig = in_cont = in_recp = in_att = False
    sig_lines = []; i = 0

    while i < len(lines):
        s = lines[i].strip()
        if not s: i += 1; continue
        s = _strip_markdown(s)

        # 特殊标记
        if s == '---RECIPIENTS---': in_recp = True; i += 1; continue
        if in_recp: result['recipients'] = s; in_recp = False; i += 1; continue
        if s == '---SIGNATURE---': in_sig = True; i += 1; continue
        if s == '---CONTACT---': in_sig = False; in_cont = True; i += 1; continue
        if in_cont:
            result['contact'] = s.strip('\uff08\uff09()'); in_cont = False; i += 1; continue
        if in_sig: sig_lines.append(s); i += 1; continue

        # 标题层级
        if s.startswith('# ') and not s.startswith('## '):
            result['title'] = s[2:]; i += 1; continue
        if s.startswith('## ') and not s.startswith('### '):
            result['body_lines'].append({"type":"para","text":s[3:],"level":1}); i += 1; continue
        if s.startswith('### ') and not s.startswith('#### '):
            result['body_lines'].append({"type":"para","text":s[4:],"level":2}); i += 1; continue
        if s.startswith('#### '):
            result['body_lines'].append({"type":"para","text":s[5:],"level":3}); i += 1; continue

        # 附件说明
        if s.startswith('\u9644\u4ef6') and ('\uff1a' in s or ':' in s or len(s)<=4):
            in_att = True; result['attachments'].append(s); i += 1; continue
        if in_att and _RE_ATTACH_NUM.match(s):
            result['attachments'].append(s); i += 1; continue

        # 附件表格（---ATTACHMENT---）→ 独立存储，在落款后渲染
        if s == '---ATTACHMENT---':
            i += 1; al = lines[i].strip() if i < len(lines) else '附件'; i += 1
            while i < len(lines) and not lines[i].strip(): i += 1
            td, ni = _parse_table_block(lines, i)
            if td is not None:
                td['is_attachment'] = True; td['attachment_label'] = al
                result['attachment_tables'].append(td); i = ni
            continue

        # 正文表格
        td, ni = _parse_table_block(lines, i)
        if td is not None: result['body_lines'].append(td); i = ni; in_att = False; continue

        # 普通段落
        result['body_lines'].append({"type":"para","text":s,"level":0}); i += 1

    if sig_lines:
        result['signature'] = {'org':'','date':''}
        for sl in sig_lines:
            if _RE_DATE.match(sl): result['signature']['date'] = sl
            else: result['signature']['org'] = sl
    return result

# ============ 文档构建 ============
def build_document(parsed, task_name, output_dir, doc_type=None):
    doc = Document()
    _setup_page(doc); _add_page_number(doc)

    if parsed['header']:
        _add_para(doc, parsed['header'], FONT_BIAOSONG_NAME, PT_36,
                  alignment=WD_ALIGN_PARAGRAPH.CENTER)
        _add_empty_line(doc)

    if parsed['doc_number']:
        _add_para(doc, parsed['doc_number'], FONT_SONG, PT_16,
                  alignment=WD_ALIGN_PARAGRAPH.CENTER)

    if parsed['title']:
        _add_heading_para(doc, parsed['title'], 'title')

    has_recipients = bool(parsed.get('recipients'))
    if has_recipients:
        _add_empty_line(doc)
        _add_para(doc, parsed['recipients'], FONT_SONG, PT_16,
                  alignment=WD_ALIGN_PARAGRAPH.LEFT,
                  first_indent=Pt(0))  # 主送机关顶格左对齐，不缩进

    if not has_recipients: _add_empty_line(doc)
    for item in parsed['body_lines']:
        if item.get('type') == 'table':
            _render_table(doc, item)
        else:
            text = item[0] if isinstance(item, (tuple,list)) else item.get('text','')
            level = item[1] if isinstance(item, (tuple,list)) else item.get('level',0)
            if level in (1, 2, 3):
                # 一级/二级/三级标题 → 挂 Heading 1/2/3 大纲级别（可一键生成目录）
                _add_heading_para(doc, text, level)
            else:
                _add_para(doc, text, FONT_SONG, PT_16, alignment=WD_ALIGN_PARAGRAPH.LEFT,
                          first_indent=INDENT_2CHAR)

    if parsed['attachments']:
        _add_empty_line(doc)
        for att in parsed['attachments']:
            _add_para(doc, att, FONT_SONG, PT_16, alignment=WD_ALIGN_PARAGRAPH.LEFT,
                      first_indent=INDENT_2CHAR)

    if parsed['signature']:
        _add_empty_line(doc, count=3)
        sig = parsed['signature']
        if sig.get('org'):
            _add_para(doc, sig['org'], FONT_SONG, PT_16, alignment=WD_ALIGN_PARAGRAPH.RIGHT)
        if sig.get('date'):
            _add_para(doc, sig['date'], FONT_SONG, PT_16, alignment=WD_ALIGN_PARAGRAPH.RIGHT)

    contact_text = parsed.get('contact') or ''
    if contact_text:
        _add_empty_line(doc)
        _add_para(doc, f"\uff08\u8054\u7cfb\u4eba\uff1a{contact_text}\uff09",
                  FONT_SONG, PT_16, alignment=WD_ALIGN_PARAGRAPH.CENTER)

    # 附件表格（另面编排，在落款之后）
    for at in parsed.get('attachment_tables', []):
        _render_table(doc, at)

    date_str = datetime.now().strftime("%Y.%m.%d")
    safe_name = task_name.replace("/","-").replace("\\","-").replace(":","\uff1a")
    filename = f"{date_str} {safe_name}.docx"
    output_path = os.path.join(output_dir, filename)
    os.makedirs(output_dir, exist_ok=True)
    doc.save(output_path)
    _embed_fonts(output_path)
    print(f"[OK] {output_path}")
    return output_path

# ============ 入口 ============
def main():
    p = argparse.ArgumentParser(description="国标公文格式化导出 v18.5")
    p.add_argument("--content","-c", help="Markdown \u5185\u5bb9")
    p.add_argument("--markdown-file","-f", help="Markdown \u6587\u4ef6\u8def\u5f84")
    p.add_argument("--task-name","-n", required=True, help="\u4efb\u52a1\u540d\u79f0")
    p.add_argument("--output","-o", default=_DEFAULT_OUTPUT,
                   help=f"\u8f93\u51fa\u76ee\u5f55\uff08\u9ed8\u8ba4\uff1a{_DEFAULT_OUTPUT}\uff09")
    p.add_argument("--doc-type","-t", help="\u6587\u79cd")
    p.add_argument("--doc-number", help="\u53d1\u6587\u5b57\u53f7")
    p.add_argument("--sender", help="\u53d1\u6587\u673a\u5173")
    p.add_argument("--recipients", help="\u4e3b\u9001\u673a\u5173")
    p.add_argument("--contact", help="\u8054\u7cfb\u4eba")
    p.add_argument("--rm-input", action="store_true",
                   help="\u5bfc\u51fa\u540e\u81ea\u52a8\u5220\u9664\u8f93\u5165md\u6587\u4ef6")
    args = p.parse_args()

    if args.markdown_file and os.path.isfile(args.markdown_file):
        with open(args.markdown_file,'r',encoding='utf-8') as f: content = f.read()
    elif args.content: content = args.content
    else: content = sys.stdin.read()

    if not content or not content.strip():
        print("[ERROR] \u5185\u5bb9\u4e3a\u7a7a"); sys.exit(1)

    parsed = parse_markdown(content, doc_type=args.doc_type, doc_number=args.doc_number,
                            sender=args.sender, recipients=args.recipients, contact=args.contact)
    build_document(parsed, args.task_name, args.output, doc_type=args.doc_type)

    if args.rm_input and args.markdown_file and os.path.isfile(args.markdown_file):
        try: os.remove(args.markdown_file)
        except OSError: pass

if __name__ == "__main__": main()
