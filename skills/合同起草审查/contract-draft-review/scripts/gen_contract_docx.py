#!/usr/bin/env python3
"""
gen_contract_docx.py — 合同 Word 文档生成脚本

支持两种模式：
  - draft: 起草模式，生成纯文本合同
  - review: 审查模式，生成带修订标记的合同（删除=红色删除线，新增=蓝色下划线）

用法:
  python3 gen_contract_docx.py <input_json_path>

输入: JSON 文件路径（格式见 SKILL.md 中的定义）
输出: 同目录下生成同名 .docx 文件
"""

import json
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml


# ── 样式常量 ──────────────────────────────────────────────

FONT_CN = "宋体"
FONT_HEADING = "黑体"
COLOR_DELETE = RGBColor(0xCC, 0x00, 0x00)   # 深红
COLOR_INSERT = RGBColor(0x00, 0x56, 0xB3)   # 深蓝
COLOR_NORMAL = RGBColor(0x00, 0x00, 0x00)   # 黑色


def set_run_font(run, name=FONT_CN, size=Pt(12), color=None, bold=False,
                 strike=False, underline=False):
    """统一设置 run 的字体、大小、颜色、加粗、删除线、下划线。"""
    run.font.size = size
    run.font.name = name
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if strike:
        run.font.strike = True
    if underline:
        run.font.underline = True
    # 设置中文字体
    r = run._element
    rPr = r.find(qn('w:rPr'))
    if rPr is None:
        rPr = parse_xml(f'<w:rPr {nsdecls("w")}></w:rPr>')
        r.insert(0, rPr)
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")}></w:rFonts>')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), name)


def set_paragraph_format(para, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
                         space_before=Pt(6), space_after=Pt(6),
                         line_spacing=1.5, first_line_indent=None):
    """统一设置段落格式。"""
    para.alignment = alignment
    pf = para.paragraph_format
    pf.space_before = space_before
    pf.space_after = space_after
    pf.line_spacing = line_spacing
    if first_line_indent is not None:
        pf.first_line_indent = first_line_indent


# ── 起草模式 ──────────────────────────────────────────────

def generate_draft(data: dict, output_path: Path):
    """起草模式：生成纯文本合同 Word 文档。"""
    doc = Document()

    # 页面设置
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(3.17)

    # 合同标题
    title = data.get("title", "合同")
    title_para = doc.add_paragraph()
    set_paragraph_format(title_para, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                         space_before=Pt(24), space_after=Pt(24))
    run = title_para.add_run(title)
    set_run_font(run, name=FONT_HEADING, size=Pt(22), bold=True)

    # 当事人信息
    parties = data.get("parties", {})
    if parties:
        for key, label in [("party_a", "甲方"), ("party_b", "乙方")]:
            val = parties.get(key, "")
            if val:
                p = doc.add_paragraph()
                set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.LEFT,
                                     first_line_indent=Cm(0.74))
                run = p.add_run(f"{label}：{val}")
                set_run_font(run, bold=True)

    # 合同正文
    for sec in data.get("sections", []):
        heading = sec.get("heading", "")
        if heading:
            h_para = doc.add_paragraph()
            set_paragraph_format(h_para, alignment=WD_ALIGN_PARAGRAPH.LEFT,
                                 space_before=Pt(12))
            run = h_para.add_run(heading)
            set_run_font(run, name=FONT_HEADING, size=Pt(14), bold=True)

        for para_data in sec.get("paragraphs", []):
            if isinstance(para_data, str):
                text = para_data
            else:
                text = para_data.get("text", "")

            if not text:
                continue

            p = doc.add_paragraph()
            set_paragraph_format(p, first_line_indent=Cm(0.74))
            run = p.add_run(text)
            set_run_font(run)

    # 落款区
    doc.add_paragraph()  # 空行
    doc.add_paragraph()
    sign_para_a = doc.add_paragraph()
    set_paragraph_format(sign_para_a, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    run = sign_para_a.add_run("甲方（签章）：________________")
    set_run_font(run)

    sign_para_b = doc.add_paragraph()
    set_paragraph_format(sign_para_b, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    run = sign_para_b.add_run("乙方（签章）：________________")
    set_run_font(run)

    date_para = doc.add_paragraph()
    set_paragraph_format(date_para, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    run = date_para.add_run("签署日期：____年____月____日")
    set_run_font(run)

    doc.save(str(output_path))


# ── 审查模式 ──────────────────────────────────────────────

def add_revision_run(para, text, rev_type):
    """在段落中添加一个修订标记的 run。

    rev_type:
      - "unchanged": 正常文本（黑色）
      - "delete": 删除文本（红色 + 删除线）
      - "insert": 新增文本（蓝色 + 下划线）
    """
    if not text:
        return

    if rev_type == "unchanged":
        run = para.add_run(text)
        set_run_font(run, color=COLOR_NORMAL)
    elif rev_type == "delete":
        run = para.add_run(text)
        set_run_font(run, color=COLOR_DELETE, strike=True)
    elif rev_type == "insert":
        run = para.add_run(text)
        set_run_font(run, color=COLOR_INSERT, underline=True)
    else:
        run = para.add_run(text)
        set_run_font(run)


def generate_review(data: dict, output_path: Path):
    """审查模式：生成带修订标记的合同 Word 文档。"""
    doc = Document()

    # 页面设置
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(3.17)

    # ── 审查摘要页 ──
    summary = data.get("review_summary", {})
    if summary:
        # 摘要标题
        h = doc.add_paragraph()
        set_paragraph_format(h, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                             space_before=Pt(24), space_after=Pt(18))
        run = h.add_run("合同审查报告")
        set_run_font(run, name=FONT_HEADING, size=Pt(22), bold=True)

        # 风险概览表
        table = doc.add_table(rows=5, cols=2, style='Table Grid')
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_data = [
            ("合同名称", data.get("title", "—")),
            ("整体风险评级", summary.get("overall_risk", "—")),
            ("问题总数", str(summary.get("total_issues", 0))),
            ("高风险", f"🔴 {summary.get('high_risk', 0)} 项"),
            ("中/低风险", f"🟡 {summary.get('medium_risk', 0)} 项 / 🟢 {summary.get('low_risk', 0)} 项"),
        ]
        for i, (label, value) in enumerate(table_data):
            cell_label = table.cell(i, 0)
            cell_value = table.cell(i, 1)
            cell_label.width = Cm(4)
            # 清空默认段落并写入
            cell_label.paragraphs[0].clear()
            run_l = cell_label.paragraphs[0].add_run(label)
            set_run_font(run_l, bold=True)
            cell_value.paragraphs[0].clear()
            run_v = cell_value.paragraphs[0].add_run(value)
            set_run_font(run_v)

        # 修订说明
        doc.add_paragraph()
        legend_para = doc.add_paragraph()
        set_paragraph_format(legend_para)
        run = legend_para.add_run("修订标记说明：")
        set_run_font(run, bold=True)
        legend_items = [
            ("删除内容", "delete"),
            ("新增内容", "insert"),
            ("，保留内容为黑色正文。", "unchanged"),
        ]
        for text, rtype in legend_items:
            add_revision_run(legend_para, text, rtype)

        # 分页
        doc.add_page_break()

    # ── 带修订标记的合同正文 ──
    title = data.get("title", "合同")
    title_para = doc.add_paragraph()
    set_paragraph_format(title_para, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                         space_before=Pt(24), space_after=Pt(24))
    run = title_para.add_run(title)
    set_run_font(run, name=FONT_HEADING, size=Pt(22), bold=True)

    for sec in data.get("sections", []):
        heading = sec.get("heading", "")
        if heading:
            h_para = doc.add_paragraph()
            set_paragraph_format(h_para, alignment=WD_ALIGN_PARAGRAPH.LEFT,
                                 space_before=Pt(12))
            run = h_para.add_run(heading)
            set_run_font(run, name=FONT_HEADING, size=Pt(14), bold=True)

        for para_data in sec.get("paragraphs", []):
            if isinstance(para_data, str):
                text, rev_type = para_data, "unchanged"
            else:
                text = para_data.get("text", "")
                rev_type = para_data.get("type", "unchanged")

            if not text:
                continue

            p = doc.add_paragraph()
            set_paragraph_format(p, first_line_indent=Cm(0.74))
            add_revision_run(p, text, rev_type)

    # ── 审查意见明细 ──
    review_notes = data.get("review_notes", [])
    if review_notes:
        doc.add_page_break()

        h = doc.add_paragraph()
        set_paragraph_format(h, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                             space_after=Pt(18))
        run = h.add_run("审查意见明细")
        set_run_font(run, name=FONT_HEADING, size=Pt(18), bold=True)

        # 审查意见表
        note_table = doc.add_table(
            rows=1 + len(review_notes),
            cols=4,
            style='Table Grid'
        )
        note_table.alignment = WD_TABLE_ALIGNMENT.CENTER

        # 表头
        headers = ["条款位置", "风险等级", "问题描述", "修改建议"]
        for j, header in enumerate(headers):
            cell = note_table.cell(0, j)
            cell.paragraphs[0].clear()
            run = cell.paragraphs[0].add_run(header)
            set_run_font(run, bold=True, size=Pt(10))

        # 数据行
        for i, note in enumerate(review_notes, 1):
            risk = note.get("risk_level", "low")
            risk_map = {"high": "🔴 高", "medium": "🟡 中", "low": "🟢 低"}
            row_data = [
                note.get("location", "—"),
                risk_map.get(risk, risk),
                note.get("issue", "—"),
                note.get("suggestion", "—"),
            ]
            for j, val in enumerate(row_data):
                cell = note_table.cell(i, j)
                cell.paragraphs[0].clear()
                run = cell.paragraphs[0].add_run(val)
                set_run_font(run, size=Pt(10))

    doc.save(str(output_path))


# ── 主入口 ────────────────────────────────────────────────

def main():
    if len(sys.argv) < 2:
        print("用法: python3 gen_contract_docx.py <input_json_path>")
        sys.exit(1)

    json_path = Path(sys.argv[1])
    if not json_path.exists():
        print(f"错误: 文件不存在: {json_path}")
        sys.exit(1)

    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    mode = data.get("mode", "draft")
    output_path = json_path.with_suffix(".docx")

    if mode == "draft":
        generate_draft(data, output_path)
    elif mode == "review":
        generate_review(data, output_path)
    else:
        print(f"错误: 未知模式 '{mode}'，支持 'draft' 或 'review'")
        sys.exit(1)

    print(f"✅ 文档已生成: {output_path}")


if __name__ == "__main__":
    main()
