#!/usr/bin/env python3
"""
gen_contract_docx.py — 合同 Word 文档生成脚本

支持两种模式：
  - draft: 起草模式，生成合同 Word 文档
  - review: 审查模式，生成带真实 Word 修订标记（Track Changes）的合同

修订标记使用 OpenXML 的 w:ins / w:del 元素实现，在 Word 中可接受/拒绝修订。

用法:
  python3 gen_contract_docx.py <input_json_path>

输入: JSON 文件路径（格式见 SKILL.md 中的定义）
输出: 同目录下生成同名 .docx 文件
"""

import json
import sys
from datetime import datetime
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml


# ── 样式常量 ──────────────────────────────────────────────

FONT_CN = "宋体"
FONT_HEADING = "黑体"
REVISION_AUTHOR = "合同审查"
REVISION_DATE = datetime.now().strftime("%Y-%m-%dT%H:%M:%SZ")

# 修订 ID 计数器（全局递增，确保每个修订标记 id 唯一）
_rev_id_counter = [0]


def next_rev_id():
    _rev_id_counter[0] += 1
    return str(_rev_id_counter[0])


def set_run_font(run, name=FONT_CN, size=Pt(12), color=None, bold=False):
    """设置 run 的字体、大小、颜色、加粗。"""
    run.font.size = size
    run.font.name = name
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    # 中文字体
    r = run._element
    rPr = r.find(qn('w:rPr'))
    if rPr is None:
        rPr = parse_xml(f'<w:rPr {nsdecls("w")}></w:rPr>')
        r.insert(0, rPr)
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")}></w:rFonts>')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), name)


def set_paragraph_format(para, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
                         space_before=Pt(6), space_after=Pt(6),
                         line_spacing=1.5, first_line_indent=None):
    """统一设置段落格式。"""
    para.alignment = alignment
    pf = para.paragraph_format
    pf.space_before = space_before
    pf.space_after = space_after
    pf.line_spacing = line_spacing
    if first_line_indent is not None:
        pf.first_line_indent = first_line_indent


def _make_run_xml(text, font_cn=FONT_CN, font_size_half_pt=24, bold=False):
    """构造一个 w:r 元素的 XML 字符串（用于嵌入 w:ins/w:del）。"""
    bold_tag = '<w:b/>' if bold else ''
    return (
        f'<w:r {nsdecls("w")}>'
        f'<w:rPr>{bold_tag}'
        f'<w:rFonts w:eastAsia="{font_cn}"/>'
        f'<w:sz w:val="{font_size_half_pt}"/>'
        f'</w:rPr>'
        f'<w:t xml:space="preserve">{_escape_xml(text)}</w:t>'
        f'</w:r>'
    )


def _make_del_run_xml(text, font_cn=FONT_CN, font_size_half_pt=24, bold=False):
    """构造删除修订中的 w:r 元素（使用 w:delText 而非 w:t）。"""
    bold_tag = '<w:b/>' if bold else ''
    return (
        f'<w:r {nsdecls("w")}>'
        f'<w:rPr>{bold_tag}'
        f'<w:rFonts w:eastAsia="{font_cn}"/>'
        f'<w:sz w:val="{font_size_half_pt}"/>'
        f'</w:rPr>'
        f'<w:delText xml:space="preserve">{_escape_xml(text)}</w:delText>'
        f'</w:r>'
    )


def _escape_xml(text):
    """转义 XML 特殊字符。"""
    return (text
            .replace('&', '&amp;')
            .replace('<', '&lt;')
            .replace('>', '&gt;')
            .replace('"', '&quot;')
            .replace("'", '&apos;'))


# ── 真实修订标记 ──────────────────────────────────────────

def add_insert_revision(para, text, font_cn=FONT_CN, font_size=Pt(12), bold=False):
    """在段落中添加一个插入修订（w:ins）。

    在 Word 中会以彩色下划线标记，可通过"审阅→接受"纳入正文。
    """
    if not text:
        return
    rid = next_rev_id()
    size_half_pt = int(font_size.pt * 2)
    run_xml = _make_run_xml(text, font_cn, size_half_pt, bold)
    ins_xml = (
        f'<w:ins {nsdecls("w")} '
        f'w:id="{rid}" '
        f'w:author="{REVISION_AUTHOR}" '
        f'w:date="{REVISION_DATE}">'
        f'{run_xml}'
        f'</w:ins>'
    )
    ins_elem = parse_xml(ins_xml)
    para._element.append(ins_elem)


def add_delete_revision(para, text, font_cn=FONT_CN, font_size=Pt(12), bold=False):
    """在段落中添加一个删除修订（w:del）。

    在 Word 中会以删除线+淡色标记，可通过"审阅→拒绝"恢复原文。
    """
    if not text:
        return
    rid = next_rev_id()
    size_half_pt = int(font_size.pt * 2)
    run_xml = _make_del_run_xml(text, font_cn, size_half_pt, bold)
    del_xml = (
        f'<w:del {nsdecls("w")} '
        f'w:id="{rid}" '
        f'w:author="{REVISION_AUTHOR}" '
        f'w:date="{REVISION_DATE}">'
        f'{run_xml}'
        f'</w:del>'
    )
    del_elem = parse_xml(del_xml)
    para._element.append(del_elem)


# ── 起草模式 ──────────────────────────────────────────────

def generate_draft(data: dict, output_path: Path):
    """起草模式：生成合同 Word 文档。"""
    doc = Document()

    # 页面设置
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(3.17)

    # 合同标题
    title = data.get("title", "合同")
    title_para = doc.add_paragraph()
    set_paragraph_format(title_para, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                         space_before=Pt(24), space_after=Pt(24))
    run = title_para.add_run(title)
    set_run_font(run, name=FONT_HEADING, size=Pt(22), bold=True)

    # 当事人信息
    parties = data.get("parties", {})
    if parties:
        for key, label in [("party_a", "甲方"), ("party_b", "乙方")]:
            val = parties.get(key, "")
            if val:
                p = doc.add_paragraph()
                set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.LEFT,
                                     first_line_indent=Cm(0.74))
                run = p.add_run(f"{label}：{val}")
                set_run_font(run, bold=True)

    # 合同正文
    for sec in data.get("sections", []):
        heading = sec.get("heading", "")
        if heading:
            h_para = doc.add_paragraph()
            set_paragraph_format(h_para, alignment=WD_ALIGN_PARAGRAPH.LEFT,
                                 space_before=Pt(12))
            run = h_para.add_run(heading)
            set_run_font(run, name=FONT_HEADING, size=Pt(14), bold=True)

        for para_data in sec.get("paragraphs", []):
            if isinstance(para_data, str):
                text = para_data
                rev_type = "unchanged"
            else:
                text = para_data.get("text", "")
                rev_type = para_data.get("type", "unchanged")

            if not text:
                continue

            p = doc.add_paragraph()
            set_paragraph_format(p, first_line_indent=Cm(0.74))

            # 起草模式下也支持修订标记（用于范本对比调整）
            if rev_type == "unchanged":
                run = p.add_run(text)
                set_run_font(run)
            elif rev_type == "insert":
                add_insert_revision(p, text)
            elif rev_type == "delete":
                add_delete_revision(p, text)

    # 落款区
    doc.add_paragraph()
    doc.add_paragraph()
    sign_para_a = doc.add_paragraph()
    set_paragraph_format(sign_para_a, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    run = sign_para_a.add_run("甲方（签章）：________________")
    set_run_font(run)

    sign_para_b = doc.add_paragraph()
    set_paragraph_format(sign_para_b, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    run = sign_para_b.add_run("乙方（签章）：________________")
    set_run_font(run)

    date_para = doc.add_paragraph()
    set_paragraph_format(date_para, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    run = date_para.add_run("签署日期：____年____月____日")
    set_run_font(run)

    return doc


# ── 审查模式 ──────────────────────────────────────────────

def generate_review(data: dict, output_path: Path):
    """审查模式：生成带真实 Word 修订标记的合同 Word 文档。"""
    doc = Document()

    # 页面设置
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(3.17)

    # ── 审查摘要页 ──
    summary = data.get("review_summary", {})
    if summary:
        h = doc.add_paragraph()
        set_paragraph_format(h, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                             space_before=Pt(24), space_after=Pt(18))
        run = h.add_run("合同审查报告")
        set_run_font(run, name=FONT_HEADING, size=Pt(22), bold=True)

        # 风险概览表
        table = doc.add_table(rows=5, cols=2, style='Table Grid')
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_data = [
            ("合同名称", data.get("title", "—")),
            ("整体风险评级", summary.get("overall_risk", "—")),
            ("问题总数", str(summary.get("total_issues", 0))),
            ("高风险", f"🔴 {summary.get('high_risk', 0)} 项"),
            ("中/低风险", f"🟡 {summary.get('medium_risk', 0)} 项 / 🟢 {summary.get('low_risk', 0)} 项"),
        ]
        for i, (label, value) in enumerate(table_data):
            cell_label = table.cell(i, 0)
            cell_value = table.cell(i, 1)
            cell_label.width = Cm(4)
            cell_label.paragraphs[0].clear()
            run_l = cell_label.paragraphs[0].add_run(label)
            set_run_font(run_l, bold=True)
            cell_value.paragraphs[0].clear()
            run_v = cell_value.paragraphs[0].add_run(value)
            set_run_font(run_v)

        # 修订说明
        doc.add_paragraph()
        legend_para = doc.add_paragraph()
        set_paragraph_format(legend_para)
        run = legend_para.add_run("修订标记说明：")
        set_run_font(run, bold=True)
        run = legend_para.add_run("本文档使用 Word 修订模式（Track Changes），在 Word 中可通过审阅选项卡接受或拒绝修订。")
        set_run_font(run)

        # 分页
        doc.add_page_break()

    # ── 带修订标记的合同正文 ──
    title = data.get("title", "合同")
    title_para = doc.add_paragraph()
    set_paragraph_format(title_para, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                         space_before=Pt(24), space_after=Pt(24))
    run = title_para.add_run(title)
    set_run_font(run, name=FONT_HEADING, size=Pt(22), bold=True)

    for sec in data.get("sections", []):
        heading = sec.get("heading", "")
        if heading:
            h_para = doc.add_paragraph()
            set_paragraph_format(h_para, alignment=WD_ALIGN_PARAGRAPH.LEFT,
                                 space_before=Pt(12))
            run = h_para.add_run(heading)
            set_run_font(run, name=FONT_HEADING, size=Pt(14), bold=True)

        for para_data in sec.get("paragraphs", []):
            if isinstance(para_data, str):
                text, rev_type = para_data, "unchanged"
            else:
                text = para_data.get("text", "")
                rev_type = para_data.get("type", "unchanged")

            if not text:
                continue

            p = doc.add_paragraph()
            set_paragraph_format(p, first_line_indent=Cm(0.74))

            if rev_type == "unchanged":
                run = p.add_run(text)
                set_run_font(run)
            elif rev_type == "insert":
                add_insert_revision(p, text)
            elif rev_type == "delete":
                add_delete_revision(p, text)

    # ── 审查意见明细 ──
    review_notes = data.get("review_notes", [])
    if review_notes:
        doc.add_page_break()

        h = doc.add_paragraph()
        set_paragraph_format(h, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                             space_after=Pt(18))
        run = h.add_run("审查意见明细")
        set_run_font(run, name=FONT_HEADING, size=Pt(18), bold=True)

        note_table = doc.add_table(
            rows=1 + len(review_notes),
            cols=4,
            style='Table Grid'
        )
        note_table.alignment = WD_TABLE_ALIGNMENT.CENTER

        headers = ["条款位置", "风险等级", "问题描述", "修改建议"]
        for j, header in enumerate(headers):
            cell = note_table.cell(0, j)
            cell.paragraphs[0].clear()
            run = cell.paragraphs[0].add_run(header)
            set_run_font(run, bold=True, size=Pt(10))

        for i, note in enumerate(review_notes, 1):
            risk = note.get("risk_level", "low")
            risk_map = {"high": "🔴 高", "medium": "🟡 中", "low": "🟢 低"}
            row_data = [
                note.get("location", "—"),
                risk_map.get(risk, risk),
                note.get("issue", "—"),
                note.get("suggestion", "—"),
            ]
            for j, val in enumerate(row_data):
                cell = note_table.cell(i, j)
                cell.paragraphs[0].clear()
                run = cell.paragraphs[0].add_run(val)
                set_run_font(run, size=Pt(10))

    doc.save(str(output_path))

    return doc


# ── 主入口 ────────────────────────────────────────────────

def highlight_pending(doc):
    """v1.1.0：将段落/表格中含【待核 的 run 加黄色底纹（FFFF00），某负责人硬性要求"""
    from lxml import etree
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    count = 0
    def mark_run(run):
        nonlocal count
        if "【待核" in (run.text or ""):
            rPr = run._element.get_or_add_rPr()
            shd = rPr.find(f"{W}shd")
            if shd is None:
                shd = etree.SubElement(rPr, f"{W}shd")
            shd.set(f"{W}val", "clear"); shd.set(f"{W}color", "auto"); shd.set(f"{W}fill", "FFFF00")
            count += 1
    for para in doc.paragraphs:
        for run in para.runs:
            mark_run(run)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        mark_run(run)
    return count


def main():
    if len(sys.argv) < 2:
        print("用法: python3 gen_contract_docx.py <input_json_path>")
        sys.exit(1)

    json_path = Path(sys.argv[1])
    if not json_path.exists():
        print(f"错误: 文件不存在: {json_path}")
        sys.exit(1)

    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    mode = data.get("mode", "draft")
    output_path = json_path.with_suffix(".docx")

    doc = None
    if mode == "draft":
        doc = generate_draft(data, output_path)
    elif mode == "review":
        doc = generate_review(data, output_path)
    else:
        print(f"错误: 未知模式 '{mode}'，支持 'draft' 或 'review'")
        sys.exit(1)

    # v1.1.0：【待核】标黄
    hl = highlight_pending(doc)
    doc.save(output_path)
    print(f"✅ 文档已生成: {output_path}" + (f"（【待核】标黄 {hl} 处）" if hl else ""))


if __name__ == "__main__":
    main()
