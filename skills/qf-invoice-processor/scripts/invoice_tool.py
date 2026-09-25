#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
invoice_tool — 本地发票/票据识别工具链（离线，零外部 API）

子命令：
  ocr   批量识别 发票/票据（PDF/JPG/PNG）→ 结构化 JSON（含 OCR 原文）
  excel 把 JSON 结果导出为格式化 Excel

依赖：rapidocr-onnxruntime（OCR）、PyMuPDF(fitz)（PDF）、Pillow、openpyxl、numpy
用法：
  python3 invoice_tool.py ocr  <输入文件或目录> -o invoice_results.json
  python3 invoice_tool.py excel <invoice_results.json> -o invoice_results.xlsx
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
from datetime import datetime
from pathlib import Path

IMG_EXT = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".webp"}
PDF_EXT = {".pdf"}

# ---------------- OCR ----------------
_engine = None


def get_engine():
    global _engine
    if _engine is None:
        from rapidocr_onnxruntime import RapidOCR
        _engine = RapidOCR()
    return _engine


def _ocr_image(path: str) -> tuple[str, list]:
    """返回 (拼接文本, 行列表)，行列表为 [{text, box}]。"""
    engine = get_engine()
    res, _ = engine(path)
    lines = []
    if res:
        for item in res:
            box, text = item[0], item[1]
            lines.append({"text": str(text), "box": box})
    text = "\n".join(l["text"] for l in lines)
    return text, lines


def _pdf_to_images(pdf_path: str, dpi: int = 200):
    import fitz  # PyMuPDF
    import tempfile

    tmpdir = tempfile.mkdtemp(prefix="inv_")
    paths = []
    doc = fitz.open(pdf_path)
    for i, page in enumerate(doc):
        pix = page.get_pixmap(dpi=dpi)
        p = os.path.join(tmpdir, f"{Path(pdf_path).stem}_p{i+1}.png")
        pix.save(p)
        paths.append(p)
    doc.close()
    return paths


def _pdf_text_layer(pdf_path: str) -> str:
    """提取 PDF 文字层文本（电子发票/数电票自带文字层，优先使用）。无文字层返回空串。"""
    try:
        import fitz
        doc = fitz.open(pdf_path)
        txt = "\n".join(page.get_text(sort=True) for page in doc)
        doc.close()
        return txt.strip()
    except Exception:
        return ""


# ---------------- 字段抽取（增值税发票为主，尽力而为） ----------------
def _clean(s: str) -> str:
    return re.sub(r"\s+", "", s or "")


def parse_fields(text: str) -> dict:
    t = text
    tc = _clean(t)  # 去空白用于紧匹配
    f = {
        "发票类型": None,
        "发票代码": None,
        "发票号码": None,
        "开票日期": None,
        "购买方名称": None,
        "购买方税号": None,
        "销售方名称": None,
        "销售方税号": None,
        "金额不含税": None,
        "税额": None,
        "价税合计": None,
    }
    # 发票类型（核心：先判"专用 / 普通"票种，兼顾电子发票/数电票）
    tcn = tc.replace("（", "(").replace("）", ")")
    for kw, name in (
        ("电子发票(专用发票)", "电子发票（专用发票）"),
        ("电子发票(普通发票)", "电子发票（普通发票）"),
        ("增值税电子专用发票", "增值税电子专用发票"),
        ("增值税电子普通发票", "增值税电子普通发票"),
        ("增值税专用发票", "增值税专用发票"),
        ("增值税普通发票", "增值税普通发票"),
        ("机动车销售统一发票", "机动车销售统一发票"),
        ("二手车销售统一发票", "二手车销售统一发票"),
        ("区块链电子发票", "区块链电子发票"),
        ("专用发票", "增值税专用发票"),
        ("普通发票", "增值税普通发票"),
    ):
        if kw in tcn:
            f["发票类型"] = name
            break
    if not f["发票类型"]:
        f["发票类型"] = "其他发票（票种未识别）" if ("发票" in tcn or "增值税" in tcn) else None

    m = re.search(r"发票代码[::：]?\s*(\d{10,12})", tc)
    if m:
        f["发票代码"] = m.group(1)
    m = re.search(r"发票号码[::：]?\s*(\d{8,20})", tc)
    if m:
        f["发票号码"] = m.group(1)
    m = re.search(r"开票日期[::：]?\s*(\d{4})年(\d{1,2})月(\d{1,2})日", tc)
    if m:
        f["开票日期"] = f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"

    # 纳税人识别号（按出现顺序：购买方在前，销售方在后）
    tin = re.findall(r"(?:纳税人识别号|统一社会信用代码)[::：]?\s*([0-9A-Z]{15,20})", tc)
    if len(tin) >= 1:
        f["购买方税号"] = tin[0]
    if len(tin) >= 2:
        f["销售方税号"] = tin[1]

    # 名称（定向提取购买方 / 销售方，兼容"购 名称：""销 名称："及竖排"购买方信息"）
    m = re.search(r"购\s*名称[::：]?\s*(.+?)(?=销\s*名称|纳税人识别号|统一社会信用代码|$)", tc)
    if not m:
        m = re.search(r"购买方名称[::：]?\s*(.+?)(?=销售方名称|纳税人识别号|统一社会信用代码|$)", tc)
    if m:
        f["购买方名称"] = m.group(1)
    m = re.search(r"销\s*名称[::：]?\s*(.+?)(?=买|售方|纳税人识别号|统一社会信用代码|收款人|复核人|开票人|地址|开户行|电话|$)", tc)
    if not m:
        m = re.search(r"销售方名称[::：]?\s*(.+?)(?=纳税人识别号|统一社会信用代码|$)", tc)
    if m:
        f["销售方名称"] = m.group(1)

    # 金额（不含税）
    m = re.search(r"(?:金额|合计金额)[::：]?[¥￥]?([\d,]+\.\d{2})", tc)
    if m:
        f["金额不含税"] = m.group(1).replace(",", "")
    # 税额
    m = re.search(r"税额[::：]?[¥￥]?([\d,]+\.\d{2})", tc)
    if m:
        f["税额"] = m.group(1).replace(",", "")
    # 价税合计
    m = re.search(r"[（(]小写[)）][::：]?[¥￥]?([\d,]+\.\d{2})", tc)
    if not m:
        m = re.search(r"价税合计[::：]?[¥￥]?\(?小写\)?[::：]?[¥￥]?([\d,]+\.\d{2})", tc)
    if m:
        f["价税合计"] = m.group(1).replace(",", "")

    # 兜底①：电子发票票面"合计 … ¥金额 ¥税额"，取相邻的"¥金额 ¥税额"对
    if not f["金额不含税"] or not f["税额"]:
        m2 = re.search(r"[¥￥]\s*([\d,]+\.\d{2})\s*[¥￥]\s*([\d,]+\.\d{2})", t)
        if m2:
            if not f["金额不含税"]:
                f["金额不含税"] = m2.group(1).replace(",", "")
            if not f["税额"]:
                f["税额"] = m2.group(2).replace(",", "")

    # 兜底：若价税合计缺失但有金额+税额，则相加
    if not f["价税合计"] and f["金额不含税"] and f["税额"]:
        try:
            f["价税合计"] = f"{float(f['金额不含税']) + float(f['税额']):.2f}"
        except Exception:
            pass
    return f


def collect_inputs(path: str) -> list[str]:
    p = Path(path)
    files = []
    if p.is_dir():
        for f in sorted(p.rglob("*")):
            if f.suffix.lower() in IMG_EXT | PDF_EXT:
                files.append(str(f))
    elif p.exists():
        files.append(str(p))
    return files


def cmd_ocr(args) -> int:
    files = collect_inputs(args.input)
    if not files:
        print(json.dumps({"error": "未找到可处理的发票文件"}, ensure_ascii=False))
        return 1
    results = []
    for fp in files:
        rec = {"file": fp, "ok": False, "source": None, "pages": [], "text": "", "fields": {}, "error": None}
        try:
            targets = []
            if Path(fp).suffix.lower() in PDF_EXT:
                # 电子发票/数电票 PDF 自带文字层，优先直接提取（准确率远高于 OCR）
                tl = _pdf_text_layer(fp)
                if len(re.sub(r"\s+", "", tl)) >= 40:
                    rec["source"] = "text-layer"
                    rec["text"] = tl
                    rec["pages"] = [{"image": os.path.basename(fp), "lines": [], "note": "pdf-text-layer"}]
                    rec["fields"] = parse_fields(tl)
                    rec["ok"] = True
                    results.append(rec)
                    continue
                targets = _pdf_to_images(fp, dpi=args.dpi)
            else:
                targets = [fp]
            all_text = []
            for t in targets:
                text, lines = _ocr_image(t)
                rec["pages"].append({"image": os.path.basename(t), "lines": lines})
                all_text.append(text)
            rec["text"] = "\n".join(all_text)
            rec["source"] = "ocr"
            rec["fields"] = parse_fields(rec["text"])
            rec["ok"] = True
        except Exception as e:  # noqa: BLE001
            rec["error"] = f"{type(e).__name__}: {e}"
        results.append(rec)
    out = {"generated_at": datetime.now().isoformat(timespec="seconds"), "count": len(results), "results": results}
    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    Path(args.output).write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    ok = sum(1 for r in results if r["ok"])
    print(json.dumps({"status": "ok", "files": len(results), "recognized": ok, "output": args.output}, ensure_ascii=False))
    return 0


def cmd_excel(args) -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    data = json.loads(Path(args.input).read_text(encoding="utf-8"))
    recs = data.get("results", data if isinstance(data, list) else [])
    cols = [
        "文件名", "发票类型", "发票代码", "发票号码", "开票日期",
        "购买方名称", "购买方税号", "销售方名称", "销售方税号",
        "金额(不含税)", "税额", "价税合计", "状态",
    ]
    keys = ["发票类型", "发票代码", "发票号码", "开票日期", "购买方名称",
            "购买方税号", "销售方名称", "销售方税号", "金额不含税", "税额", "价税合计"]

    wb = Workbook()
    ws = wb.active
    ws.title = "发票汇总"
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="CCCCCC")
    head_font = Font(bold=True, size=11)
    align_c = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = align_c
        cell.border = border

    total = 0.0
    for r in recs:
        f = r.get("fields", {}) or {}
        row = [os.path.basename(r.get("file", ""))] + [f.get(k) for k in keys] + ["成功" if r.get("ok") else f"失败: {r.get('error')}"]
        ws.append(row)
        for c in range(1, len(cols) + 1):
            ws.cell(row=ws.max_row, column=c).border = border
        # 发票类型着色：专用发票红字、普通发票蓝字，便于一眼区分票种
        tval = f.get("发票类型") or ""
        if "专用" in tval:
            ws.cell(row=ws.max_row, column=2).font = Font(color="C00000", bold=True)
        elif "普通" in tval:
            ws.cell(row=ws.max_row, column=2).font = Font(color="1F4E79", bold=True)
        try:
            total += float(f.get("价税合计") or 0)
        except Exception:
            pass

    # 合计行
    ws.append(["合计", "", "", "", "", "", "", "", "", "", "", f"{total:.2f}", ""])
    for c in range(1, len(cols) + 1):
        ws.cell(row=ws.max_row, column=c).font = Font(bold=True)
        ws.cell(row=ws.max_row, column=c).border = border

    widths = [26, 16, 14, 18, 12, 28, 20, 28, 20, 14, 12, 14, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    print(json.dumps({"status": "ok", "rows": len(recs), "total": round(total, 2), "output": args.output}, ensure_ascii=False))
    return 0


def _preflight():
    missing = []
    for mod, pip in (("rapidocr_onnxruntime", "rapidocr-onnxruntime"),
                     ("fitz", "PyMuPDF"), ("openpyxl", "openpyxl"), ("PIL", "Pillow")):
        try:
            __import__(mod)
        except Exception:
            missing.append(pip)
    if missing:
        print("[invoice-processor] 缺少依赖：%s" % ", ".join(missing), file=sys.stderr)
        print("请安装：pip install " + " ".join(missing), file=sys.stderr)
        sys.exit(2)


def main() -> int:
    _preflight()
    p = argparse.ArgumentParser(description="本地发票/票据识别工具链（离线）")
    sub = p.add_subparsers(dest="cmd", required=True)
    p1 = sub.add_parser("ocr", help="批量识别发票 → JSON")
    p1.add_argument("input", help="输入文件或目录")
    p1.add_argument("-o", "--output", default="invoice_results.json")
    p1.add_argument("--dpi", type=int, default=200, help="PDF 转图 DPI")
    p1.set_defaults(func=cmd_ocr)
    p2 = sub.add_parser("excel", help="JSON → Excel")
    p2.add_argument("input", nargs="?", default="invoice_results.json")
    p2.add_argument("-o", "--output", default="invoice_results.xlsx")
    p2.set_defaults(func=cmd_excel)
    args = p.parse_args()
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
