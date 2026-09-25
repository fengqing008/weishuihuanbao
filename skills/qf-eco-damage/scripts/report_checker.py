#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
生态环境损害鉴定评估报告书质检脚本
五项大检查：章节完整性、占位符清理、主结论齐备、数据溯源、系数取值合规。
支持 .md / .txt / .docx（docx 需 python-docx）。

用法示例：
  python3 report_checker.py --file 报告书.md
  python3 report_checker.py --file 报告书.docx --json
退出码：0 全部通过；1 存在 FAIL。
"""
import argparse
import json
import os
import re
import sys

CHAPTERS = {
    "基本情况": ["基本情况"],
    "鉴定评估方案": ["鉴定评估方案", "鉴定评估目标"],
    "鉴定评估过程与分析": ["鉴定评估过程", "调查过程", "损害确认"],
    "损害价值量化": ["损害价值量化", "损害数额", "价值量化"],
    "恢复建议方案": ["恢复建议", "恢复方案"],
    "鉴定评估结论": ["鉴定评估结论", "结论"],
    "特别事项说明": ["特别事项说明"],
}

BAD_PLACEHOLDERS = ["TODO", "待补充", "某某", "XXX", "xxx", "（待填）", "(待填)", "???"]
TRACE_KEYWORDS = ["来源", "附件", "卷宗", "监测报告", "调查表", "踏勘", "访谈", "依据"]


def read_text(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        try:
            from docx import Document
        except ImportError:
            print("[错误] 读取 docx 需安装 python-docx：pip install python-docx", file=sys.stderr)
            sys.exit(2)
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                for c in row.cells:
                    parts.append(c.text)
        return "\n".join(parts)
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def check_chapters(text):
    missing = [name for name, kws in CHAPTERS.items() if not any(k in text for k in kws)]
    return (len(missing) == 0), missing


def check_placeholders(text):
    found = [p for p in BAD_PLACEHOLDERS if p in text]
    return (len(found) == 0), found


def check_conclusion(text):
    ok = ("鉴定评估结论" in text or "结论" in text) and ("损害" in text)
    return ok, [] if ok else ["未检出鉴定评估结论或损害结论"]


def check_trace(text):
    hits = [k for k in TRACE_KEYWORDS if k in text]
    return (len(hits) >= 2), hits


def check_coefficients(text):
    problems = []
    # 危害系数 α 取值 1~2
    for m in re.finditer(r"[αa]\s*[=＝]\s*([0-9]+(?:\.[0-9]+)?)", text):
        v = float(m.group(1))
        if not (1.0 <= v <= 2.0):
            problems.append("危害系数 α=%s 超出表1/表2 范围(1~2)" % m.group(1))
    # 超标系数 τ 取值 1~2
    for m in re.finditer(r"[τt]\s*[=＝]\s*([0-9]+(?:\.[0-9]+)?)", text):
        v = float(m.group(1))
        if not (1.0 <= v <= 2.0):
            problems.append("超标系数 τ=%s 超出表3/表4 范围(1~2)" % m.group(1))
    # 环境功能系数 ω 取值 1.5~2.5
    for m in re.finditer(r"[ωw]\s*[=＝]\s*([0-9]+(?:\.[0-9]+)?)", text):
        v = float(m.group(1))
        if not (1.5 <= v <= 2.5):
            problems.append("环境功能系数 ω=%s 超出表5 范围(1.5~2.5)" % m.group(1))
    return (len(problems) == 0), problems


def main(argv=None):
    p = argparse.ArgumentParser(description="生态环境损害鉴定评估报告书质检（五门检查）")
    p.add_argument("--file", required=True, help="待质检文件（.md/.txt/.docx）")
    p.add_argument("--json", action="store_true", help="以 JSON 输出")
    args = p.parse_args(argv)

    if not os.path.exists(args.file):
        print("[错误] 文件不存在：%s" % args.file, file=sys.stderr)
        return 2

    text = read_text(args.file)
    results = [
        ("门禁1 章节完整性",) + check_chapters(text),
        ("门禁2 占位符清理",) + check_placeholders(text),
        ("门禁3 主结论齐备",) + check_conclusion(text),
        ("门禁4 数据溯源",) + check_trace(text),
        ("门禁5 系数合规",) + check_coefficients(text),
    ]
    all_pass = all(r[1] for r in results)

    if args.json:
        print(json.dumps({
            "file": args.file,
            "pass": all_pass,
            "checks": [{"name": r[0], "pass": r[1], "detail": r[2]} for r in results],
        }, ensure_ascii=False, indent=2))
        return 0 if all_pass else 1

    print("== 报告书质检报告 ==")
    print("文件：%s" % os.path.abspath(args.file))
    print("-" * 40)
    for name, ok, detail in results:
        flag = "PASS" if ok else "FAIL"
        line = "[%s] %s" % (flag, name)
        if detail:
            line += "  -> %s" % ("；".join(str(d) for d in detail) if isinstance(detail, list) else detail)
        print(line)
    print("-" * 40)
    print("总判定：%s" % ("全部通过" if all_pass else "存在 FAIL，须回退补正"))
    return 0 if all_pass else 1


if __name__ == "__main__":
    sys.exit(main())
