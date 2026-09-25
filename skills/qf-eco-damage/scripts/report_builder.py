#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
生态环境损害鉴定评估报告书框架生成器
按 GB/T 39791.1-2020 附录A 生成七章骨架（Markdown），可同时导出 Word（标题挂 Heading 大纲级别 + 可一键更新目录域）。

用法示例：
  python3 report_builder.py --case "某某公司水污染案" --out 报告框架.docx --md 报告框架.md
  python3 report_builder.py --case "某案" --md 报告框架.md
"""
import argparse
import os
import sys

# 七章骨架：(章标题, [节标题...])
SECTIONS = [
    ("一 基本情况", ["1.1 委托方与鉴定评估机构", "1.2 案情概况", "1.3 区域环境概况"]),
    ("二 鉴定评估方案", ["2.1 鉴定评估目标", "2.2 鉴定评估依据", "2.3 鉴定评估原则",
                    "2.4 鉴定评估范围", "2.5 鉴定评估内容", "2.6 鉴定评估方法"]),
    ("三 鉴定评估过程与分析", ["3.1 调查工作内容与调查方案", "3.2 调查过程",
                        "3.3 调查结果", "3.4 环境基线确定", "3.5 损害确认",
                        "3.6 因果关系分析", "3.7 调查结论"]),
    ("四 损害价值量化", ["4.1 计算方法", "4.2 排放数量确定", "4.3 单位治理成本确定",
                   "4.4 调整系数确定", "4.5 损害数额计算与不确定性分析"]),
    ("五 生态环境损害恢复建议方案", ["5.1 恢复目标", "5.2 备选恢复方案比选",
                          "5.3 推荐恢复方案与替代性修复"]),
    ("六 鉴定评估结论", ["6.1 事实认定", "6.2 损害确定", "6.3 因果关系", "6.4 损害量化与责任建议"]),
    ("七 特别事项说明", ["7.1 报告真实性与合法性", "7.2 使用范围", "7.3 不确定性提示"]),
]

FILL = "【待填：%s】"


def build_markdown(case):
    lines = []
    lines.append("# %s生态环境损害鉴定评估报告书" % case)
    lines.append("")
    lines.append("> 依据 GB/T 39791.1—2020 附录A 编制。封面、声明、目录、落款签章、附件目录须齐备。")
    lines.append("")
    lines.append("## 声明")
    lines.append("")
    lines.append("1. 本机构及鉴定人依法独立进行鉴定评估，与本案当事人无利害关系。")
    lines.append("2. 本报告仅供委托方用于%s，采信权归委托方。" % FILL % "用途")
    lines.append("3. 本报告完整有效，复制件无效。")
    lines.append("4. 对鉴定评估意见有异议的，可依法向委托方或有关部门提出。")
    lines.append("")
    lines.append("## 目录")
    lines.append("")
    lines.append("（Word 版插入可一键更新的目录域）")
    lines.append("")
    for chap, subs in SECTIONS:
        lines.append("## %s" % chap)
        lines.append("")
        for s in subs:
            lines.append("### %s" % s)
            lines.append("")
            lines.append(FILL % s.split(" ", 1)[-1])
            lines.append("")
    lines.append("## 落款")
    lines.append("")
    lines.append("司法鉴定人（签名）：%s　执业证号：%s" % (FILL % "姓名", FILL % "证号"))
    lines.append("")
    lines.append("鉴定评估机构（盖章）：%s" % (FILL % "机构名称"))
    lines.append("")
    lines.append("出具日期：%s" % (FILL % "年月日"))
    lines.append("")
    lines.append("## 附件目录")
    lines.append("")
    for i, name in enumerate(
            ["鉴定机构及鉴定人证件", "企业地理位置图", "司法鉴定评估委托书", "司法鉴定评估告知书",
             "案件卷宗相关资料", "周边环境质量报告", "治理成本调查表", "现场踏勘记录", "人员访谈记录"], 1):
        lines.append("- 附件%s　%s" % ("一二三四五六七八九"[i - 1], name))
    lines.append("")
    return "\n".join(lines)


def add_toc(doc):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    p = doc.add_paragraph()
    r = p.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
    it.text = 'TOC \\o "1-3" \\h \\z \\u'
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "（右键→更新域，生成目录）"
    f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end")
    for e in (f1, it, f2, t, f3):
        r._r.append(e)
    try:
        settings = doc.settings.element
        uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true")
        settings.append(uf)
    except Exception:
        pass


def build_docx(case, out_path):
    try:
        from docx import Document
    except ImportError:
        print("[提示] 未安装 python-docx，跳过 Word 生成，仅输出 Markdown。可 pip install python-docx 后重试。",
              file=sys.stderr)
        return False
    doc = Document()
    doc.add_heading("%s生态环境损害鉴定评估报告书" % case, level=0)
    doc.add_heading("目录", level=1)
    add_toc(doc)
    doc.add_heading("声明", level=1)
    for line in ["本机构及鉴定人依法独立进行鉴定评估，与本案当事人无利害关系。",
                 "本报告仅供委托方使用，采信权归委托方。",
                 "本报告完整有效，复制件无效。",
                 "对鉴定评估意见有异议的，可依法向委托方或有关部门提出。"]:
        doc.add_paragraph(line)
    for chap, subs in SECTIONS:
        doc.add_heading(chap, level=1)
        for s in subs:
            doc.add_heading(s, level=2)
            doc.add_paragraph(FILL % s.split(" ", 1)[-1])
    doc.add_heading("落款", level=1)
    doc.add_paragraph("司法鉴定人（签名）：%s　执业证号：%s" % (FILL % "姓名", FILL % "证号"))
    doc.add_paragraph("鉴定评估机构（盖章）：%s" % (FILL % "机构名称"))
    doc.add_paragraph("出具日期：%s" % (FILL % "年月日"))
    doc.add_heading("附件目录", level=1)
    for i, name in enumerate(
            ["鉴定机构及鉴定人证件", "企业地理位置图", "司法鉴定评估委托书", "司法鉴定评估告知书",
             "案件卷宗相关资料", "周边环境质量报告", "治理成本调查表", "现场踏勘记录", "人员访谈记录"], 1):
        doc.add_paragraph("附件%s　%s" % ("一二三四五六七八九"[i - 1], name), style="List Bullet")
    doc.save(out_path)
    return True


def main(argv=None):
    p = argparse.ArgumentParser(description="生态环境损害鉴定评估报告书框架生成器（GB/T 39791.1-2020 附录A）")
    p.add_argument("--case", required=True, help="案件名称，如 某某公司水污染案")
    p.add_argument("--out", default=None, help="输出 Word 路径（.docx）")
    p.add_argument("--md", default=None, help="输出 Markdown 路径（.md）")
    args = p.parse_args(argv)

    md = build_markdown(args.case)
    if args.md:
        with open(args.md, "w", encoding="utf-8") as f:
            f.write(md)
        print("[生成] Markdown 框架：%s" % os.path.abspath(args.md))
    if args.out:
        ok = build_docx(args.case, args.out)
        if ok:
            print("[生成] Word 框架：%s" % os.path.abspath(args.out))
            print("[提示] 打开 Word 后按 Ctrl+A、F9 更新目录域。")
    if not args.md and not args.out:
        print(md)
    return 0


if __name__ == "__main__":
    sys.exit(main())
