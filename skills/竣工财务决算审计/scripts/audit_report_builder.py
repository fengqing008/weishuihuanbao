#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
audit_report_builder.py — 竣工财务决算审计报告生成器

按审计报告体例生成 Word：挂接 Heading 1/2 大纲级别、插入可一键更新的目录域、
字体按 GB/T 9704-2012 公文参数。支持两类体例：
  --type gov  审计机关版（正文八要素）
  --type cpa  注册会计师审核报告版（披露事项式）

用法:
  python3 scripts/audit_report_builder.py --input 报告数据.json --out 审计报告.docx --type cpa
  python3 scripts/audit_report_builder.py --demo --out 审计报告.docx --type gov

退出码:
  0 = 成功（docx 已生成）
  1 = python-docx 不可用，已降级输出 Markdown
  2 = 输入错误
"""
import argparse
import json
import os
import sys

GOV_SECTIONS = [
    ("一、审计依据、审计开展及相关责任说明", []),
    ("二、项目建设情况", []),
    ("三、审计评价意见", []),
    ("四、审计结果", []),
    ("五、审计发现问题和处理意见", []),
    ("六、审计整改及审计成效", []),
    ("七、审计建议", []),
    ("八、审计附件", []),
]

CPA_SECTIONS = [
    ("一、项目概况", ["（一）建设单位概况", "（二）建设项目概况", "（三）项目批复情况"]),
    ("二、审核依据", []),
    ("三、审核范围与程序", []),
    ("四、项目建设资金到位与使用情况", []),
    ("五、竣工财务决算报表审核情况", []),
    ("六、概算执行情况分析", []),
    ("七、交付使用资产情况", []),
    ("八、债权债务清理情况", []),
    ("九、审计（审核）发现的问题及整改建议", []),
    ("十、审核结论", []),
]


def default_data(kind):
    if kind == "cpa":
        return {
            "title": "××建设项目竣工财务决算审核报告",
            "report_no": "××审字〔20××〕××号",
            "entity": "××公司（委托单位）",
            "project": "××建设项目",
            "sections": [{"heading": h, "body": b} for h, b in CPA_SECTIONS],
            "signature": ["××会计师事务所（盖章）", "注册会计师：×××（签名）", "年  月  日"],
        }
    return {
        "title": "××建设项目竣工财务决算审计报告",
        "report_no": "×审××报〔20××〕××号",
        "entity": "××单位（被审计单位）",
        "project": "××建设项目竣工决算审计",
        "sections": [{"heading": h, "body": b} for h, b in GOV_SECTIONS],
        "signature": ["××审计机关（盖章）", "年  月  日"],
    }


def build_docx(data, out):
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    def set_font(run, ea="仿宋_GB2312", ascii_="Times New Roman", size=None, bold=None):
        run.font.name = ascii_
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        rFonts.set(qn("w:eastAsia"), ea)
        rFonts.set(qn("w:ascii"), ascii_)
        rFonts.set(qn("w:hAnsi"), ascii_)
        if size:
            run.font.size = Pt(size)
        if bold is not None:
            run.font.bold = bold

    def add_toc(doc):
        p = doc.add_paragraph()
        run = p.add_run()
        r = run._r
        f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
        it.text = 'TOC \\o "1-3" \\h \\z \\u'
        f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
        t = OxmlElement("w:t"); t.text = "（右键→更新域，生成目录）"
        f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end")
        for e in (f1, it, f2, t, f3):
            r.append(e)

    def enable_update_fields(doc):
        settings = doc.settings.element
        el = OxmlElement("w:updateFields"); el.set(qn("w:val"), "true")
        settings.append(el)

    doc = Document()

    # 标题（Title 样式，不计入目录）
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(data.get("title", "竣工财务决算审计报告"))
    set_font(run, ea="方正小标宋简体", size=22, bold=True)
    try:
        p.style = doc.styles["Title"]
    except KeyError:
        pass

    if data.get("report_no"):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(data["report_no"]), size=12)
    if data.get("entity"):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run("被审计单位（收件人）：" + data["entity"]), size=12)

    # 目录
    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("目　录"), ea="黑体", size=16, bold=True)
    add_toc(doc)
    doc.add_page_break()

    # 章节
    for sec in data.get("sections", []):
        h = doc.add_heading(sec.get("heading", ""), level=1)
        for r in h.runs:
            set_font(r, ea="黑体", size=16, bold=True)
        for line in sec.get("body", []):
            if line.startswith("（") or line.startswith("("):
                sp = doc.add_heading(line, level=2)
                for r in sp.runs:
                    set_font(r, ea="楷体_GB2312", size=16, bold=True)
            else:
                pp = doc.add_paragraph()
                pp.paragraph_format.first_line_indent = Pt(32)
                pp.paragraph_format.line_spacing = Pt(28)
                set_font(pp.add_run(line), size=16)

    # 落款
    doc.add_paragraph()
    for line in data.get("signature", []):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_font(p.add_run(line), size=16)

    enable_update_fields(doc)
    doc.save(out)
    return True


def build_markdown(data, out_md):
    lines = [f"# {data.get('title', '')}", ""]
    if data.get("report_no"):
        lines += [f"文号：{data['report_no']}", ""]
    lines += ["## 目录", ""]
    for sec in data.get("sections", []):
        lines.append(sec.get("heading", ""))
    lines.append("")
    for sec in data.get("sections", []):
        lines += [f"## {sec.get('heading', '')}", ""]
        for line in sec.get("body", []):
            lines += [f"### {line}" if line.startswith("（") else line, ""]
    for line in data.get("signature", []):
        lines.append(line)
    with open(out_md, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    return out_md


def main():
    ap = argparse.ArgumentParser(description="竣工财务决算审计报告生成器")
    ap.add_argument("--input", help="报告数据 JSON；缺省用内置骨架")
    ap.add_argument("--demo", action="store_true", help="用内置骨架生成样例")
    ap.add_argument("--out", required=True, help="输出 docx 路径")
    ap.add_argument("--type", choices=["cpa", "gov"], default="cpa", help="报告体例")
    args = ap.parse_args()

    if args.input:
        try:
            with open(args.input, encoding="utf-8") as f:
                data = json.load(f)
        except FileNotFoundError:
            print(f"错误：输入文件不存在：{args.input}", file=sys.stderr)
            sys.exit(2)
        except json.JSONDecodeError as e:
            print(f"错误：JSON 解析失败：{e}", file=sys.stderr)
            sys.exit(2)
    else:
        data = default_data(args.type)

    try:
        build_docx(data, args.out)
        print(f"审计报告已生成：{args.out}")
        sys.exit(0)
    except ImportError:
        md = os.path.splitext(args.out)[0] + ".md"
        build_markdown(data, md)
        print(f"警告：python-docx 不可用，已降级输出 Markdown：{md}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
