#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
audit_working_paper.py — 审计工作底稿生成器

按审计事项生成 Excel 工作底稿台账，含底稿编号、审计事项、审计程序与方法、
证据来源、核对结论、发现的偏差与处理、编制人、日期、复核人。
支持内置审计事项模板（六大审计维度），便于进场时快速建底稿。

用法:
  python3 scripts/audit_working_paper.py --demo --out 审计工作底稿.xlsx
  python3 scripts/audit_working_paper.py --input 底稿数据.json --out 审计工作底稿.xlsx

退出码:
  0 = 成功
  2 = 输入错误或缺少依赖
"""
import argparse
import json
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
except ImportError:
    print("错误：缺少 openpyxl，请先 pip install openpyxl", file=sys.stderr)
    sys.exit(2)

HEADERS = ["底稿编号", "审计事项", "审计程序与方法", "证据来源", "核对结论",
           "发现的偏差与处理", "编制人", "编制日期", "复核人"]

# 六大审计维度内置底稿事项模板
DIMENSION_TEMPLATE = {
    "建设程序合规性": ["立项批复", "可研批复", "初步设计批复", "概算批复",
                       "招投标合规性", "合同签订与履行", "监理制度执行", "竣工验收"],
    "财务收支": ["资金来源与到位", "资金使用与专款专用", "银行账户管理",
                 "货币资金盘点", "结余资金", "往来款项函证"],
    "工程造价": ["工程价款结算复核", "工程量抽查", "综合单价复核",
                 "工程变更与签证", "取费与措施费复核"],
    "建设管理": ["合同管理", "变更审批链", "现场管理", "内控制度执行"],
    "资产形成与移交": ["交付使用资产真实性", "资产计价", "待摊投资分摊",
                       "资产移交手续", "产权核实"],
    "投资绩效": ["概算执行偏差", "投资控制成效", "经济效益", "社会效益"],
}


def build_template_rows():
    rows = []
    idx = 1
    for dim, items in DIMENSION_TEMPLATE.items():
        for it in items:
            rows.append({
                "编号": f"WP-{idx:03d}",
                "事项": f"{dim}—{it}",
                "程序": "",
                "证据": "",
                "结论": "",
                "偏差": "",
            })
            idx += 1
    return rows


def load_rows(path):
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    rows = []
    for i, r in enumerate(data.get("items", []), start=1):
        rows.append({
            "编号": r.get("no", f"WP-{i:03d}"),
            "事项": r.get("matter", ""),
            "程序": r.get("procedure", ""),
            "证据": r.get("evidence", ""),
            "结论": r.get("conclusion", ""),
            "偏差": r.get("deviation", ""),
        })
    return rows


def write_xlsx(rows, out, project="", editors="", reviser="", date=""):
    wb = Workbook()
    ws = wb.active
    ws.title = "审计工作底稿"

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="CCCCCC")

    if project:
        ws.append([f"项目名称：{project}"])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
        ws.cell(row=1, column=1).font = Font(bold=True, size=12)
        ws.cell(row=1, column=1).alignment = Alignment(horizontal="center")
        ws.append([])

    ws.append(HEADERS)
    head_row = ws.max_row
    for c in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=head_row, column=c)
        cell.font = Font(bold=True)
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for r in rows:
        ws.append([r["编号"], r["事项"], r["程序"], r["证据"], r["结论"],
                   r["偏差"], editors, date, reviser])
        for c in range(1, len(HEADERS) + 1):
            ws.cell(row=ws.max_row, column=c).border = border
            ws.cell(row=ws.max_row, column=c).alignment = Alignment(
                vertical="top", wrap_text=True)

    widths = [12, 26, 28, 22, 26, 26, 10, 12, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    wb.save(out)
    return len(rows)


def main():
    ap = argparse.ArgumentParser(description="审计工作底稿生成器")
    ap.add_argument("--input", help="底稿数据 JSON；缺省则用内置审计事项模板")
    ap.add_argument("--demo", action="store_true", help="用内置模板生成样例底稿")
    ap.add_argument("--out", required=True, help="输出 xlsx 路径")
    ap.add_argument("--project", default="", help="项目名称")
    ap.add_argument("--editors", default="", help="编制人")
    ap.add_argument("--reviser", default="", help="复核人")
    ap.add_argument("--date", default="", help="编制日期")
    args = ap.parse_args()

    if args.input:
        try:
            rows = load_rows(args.input)
        except FileNotFoundError:
            print(f"错误：输入文件不存在：{args.input}", file=sys.stderr)
            sys.exit(2)
        except json.JSONDecodeError as e:
            print(f"错误：JSON 解析失败：{e}", file=sys.stderr)
            sys.exit(2)
    else:
        rows = build_template_rows()

    n = write_xlsx(rows, args.out, args.project, args.editors, args.reviser, args.date)
    print(f"审计工作底稿已生成：{args.out}（{n} 条底稿事项）")
    sys.exit(0)


if __name__ == "__main__":
    main()
