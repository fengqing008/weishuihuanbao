#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""报价评分助手 v2 —— 解析报价评分办法，输出得分敏感性、报价复核清单与响应条目。

边界：本脚本只做「评分规则解析与得分推演」，不代做报价测算（人材机、成本、单价分析）；
需测算报价时联动造价技能（zaojia-dinge / cost-estimator），结果回填本文第五节指定的落格。

用法:
    python3 price_scoring_helper.py --file 招标文件.md --out 报价核对表.xlsx
    python3 price_scoring_helper.py --file 招标文件.md --base 0.83 --price 0.80 --out 报价核对表.md
"""
import argparse
import re
import sys

CHECKLIST = [
    ("限价", "投标报价不超过最高投标限价（工程）／采购预算或最高限价（政府采购）",
     "超限价：工程按《招标投标法实施条例》第五十一条第（五）项否决投标；政府采购按采购文件规定的无效响应情形处理"),
    ("唯一性", "只提交一个有效报价，且报价唯一", "多方案报价未声明主选方案的，作无效处理"),
    ("成本价（工程）", "报价不低于成本",
     "低于成本报价的，评标委员会应当否决其投标（《招标投标法实施条例》第五十一条第（五）项）"),
    ("成本价（政府采购）", "不存在明显低于其他通过符合性审查投标人报价的情形",
     "报价明显低于其他通过符合性审查投标人报价、可能影响产品质量或不能诚信履约的，须在评标现场合理时间内书面说明，"
     "不能证明报价合理性的作无效投标处理（财政部令第87号第六十条）"),
    ("大小写", "报价大写与小写金额一致", "不一致且无法认定的，按不利解释处理"),
    ("一致性", "报价汇总表与分项报价表合计一致", "分项加总与总价不符即被质疑"),
    ("完整性", "报价含全部费用，未漏项、未另加条件", "漏项视为已包含在总价内"),
    ("有效期", "报价在投标有效期内固定不变", "投标有效期内不得调整报价"),
    ("电子标", "电子标报价文件与纸质／上传件完全一致", "不一致以招标文件规定为准"),
]

COST_KWS = ("低于成本", "明显低于其他投标人平均报价", "低于其他投标人", "低于其他供应商平均报价",
            "低于平均报价", "异常低价")
CN_NUM = "一二三四五六七八九十"


def lines_of(path):
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read().replace("\r\n", "\n").split("\n")


def find_price_points(lines):
    """行式通道（阿拉伯或中文序号）与表格通道（前附表/评审因素表）。"""
    pat = r"(报价分|价格分|价格评审|投标报价)[^\n。；]{0,12}?(\d{1,3}(?:\.\d)?)\s*分"
    for i, s in enumerate(lines, 1):
        if s.strip().startswith("|"):
            cells = [c.strip() for c in s.strip().strip("|").split("|")]
            if all(set(c) <= set("-: ") for c in cells):
                continue
            if any(k in s for k in ("报价分", "价格分", "价格评审", "投标报价")):
                vals = [float(m.group(1)) for c in cells
                        if (m := re.fullmatch(r"(\d{1,3}(?:\.\d{1,2})?)\s*分?", c)) and float(m.group(1)) <= 100]
                if vals:
                    return (vals[-1] if len(vals) == 1 else vals[-1]), i, s.strip()[:140]
        m = re.search(pat, s)
        if m:
            return float(m.group(2)), i, s.strip()[:140]
    return None, 0, ""


def find_rules(lines, *kws):
    out = []
    for i, s in enumerate(lines, 1):
        if any(k in s for k in kws) and not s.strip().startswith("#"):
            out.append((i, s.strip()[:200]))
    return out


def find_deduct(lines):
    out = []
    for i, s in enumerate(lines, 1):
        for m in re.finditer(r"每\s*(高|低)\s*于?[^\n。；|]{0,8}?(\d+(?:\.\d+)?)\s*%[^\n。；|]{0,8}?扣\s*(\d+(?:\.\d+)?)\s*分", s):
            out.append({"dir": m.group(1), "pct": float(m.group(2)), "ded": float(m.group(3)),
                        "line": i, "text": s.strip()[:200]})
    return out


def pick(rules):
    for r in rules:
        if any(ch.isdigit() for ch in r[1]):
            return r
    return None


def sensitivity(points, deducts):
    """按扣分规则推演得分；每条规则标注偏差方向，避免高低方向混用。"""
    if points is None or not deducts:
        return []
    rows = []
    for d in deducts:
        k = d["ded"] / d["pct"]
        for dev in (0, 1, 2, 3, 4, 5):
            score = max(points - k * dev, 0.0)
            rows.append([f"每{d['dir']}于基准价{d['pct']:g}%扣{d['ded']:g}分",
                         f"报价{d['dir']}于基准价", dev, round(score, 2), f"L{d['line']}"])
    return rows


def build_sections(path, base=None, price=None):
    lines = lines_of(path)
    pts, pts_line, _ = find_price_points(lines)
    lim = pick(find_rules(lines, "最高投标限价", "招标控制价", "最高限价"))
    bro = pick(find_rules(lines, "评标基准价"))
    cst = pick(find_rules(lines, *COST_KWS))
    mod = pick(find_rules(lines, "固定总价", "固定单价", "费率", "下浮率"))
    deducts = find_deduct(lines)
    sens = sensitivity(pts, deducts)

    summary = [["项目", "结论（机抓·待核）", "原文出处"]]
    summary.append(["报价分值", f"{pts:g} 分" if pts else "未命中（人工录入）", f"L{pts_line}" if pts else "—"])
    summary.append(["最高投标限价/控制价", lim[1] if lim else "未命中（人工录入）", f"L{lim[0]}" if lim else "—"])
    summary.append(["评标基准价规则", bro[1] if bro else "未命中（人工录入）", f"L{bro[0]}" if bro else "—"])
    if deducts:
        for d in deducts:
            summary.append(["扣分规则", f"每{d['dir']}于基准价 {d['pct']:g}% 扣 {d['ded']:g} 分", f"L{d['line']}"])
    else:
        summary.append(["扣分规则", "未命中（可能为分档制或公式制，须人工录入）", "—"])
    summary.append(["成本价/异常低价条款", cst[1] if cst else "未命中（人工录入）", f"L{cst[0]}" if cst else "—"])
    summary.append(["报价形式", mod[1] if mod else "未命中（人工录入）", f"L{mod[0]}" if mod else "—"])

    sens_rows = [["扣分规则", "偏差方向", "偏差(%)（绝对值）", "按规则推演得分", "出处"]] + sens

    demo = []
    if base is not None and price is not None and pts is not None:
        dev = (price - base) / base * 100
        demo.append(["基准价", f"{base:g}", "—", "—", "命令参数"])
        demo.append(["投标报价", f"{price:g}", f"{dev:+.2f}%（相对基准价）", "—", "命令参数"])
        if deducts:
            for d in deducts:
                k = d["ded"] / d["pct"]
                eff = abs(dev) if (dev > 0 and d["dir"] == "高") or (dev < 0 and d["dir"] == "低") else 0.0
                demo.append([f"每{d['dir']}{d['pct']:g}%扣{d['ded']:g}分", f"报价{d['dir']}于基准价",
                             f"{eff:.2f}", f"{max(pts - k * eff, 0):.2f}", f"L{d['line']}"])
        else:
            demo.append(["推演", "未解析出可计算扣分规则", "—", "须按招标文件公式手工推演", "—"])

    check_rows = [["类别", "核查项", "不满足的后果", "是否满足", "证据位置"]]
    for cat, item, cons in CHECKLIST:
        check_rows.append([cat, item, cons, "□", ""])

    resp_rows = [["序号", "评审因素(含分值)", "评分标准(分档原文)", "响应情况", "响应位置"],
                 ["—", f"投标报价（{pts:g}分·待核）" if pts else "投标报价（分值待核）",
                  "见报价评分办法摘要",
                  "完全响应：投标报价。报价不超过最高投标限价，唯一报价且大小写一致，已包含全部费用无漏项，"
                  "附报价合理性与成本说明；报价明细见《报价表》与报价说明（报价数据须先经造价核价确认后定稿）。",
                  "详见《报价表》及报价说明章节"]]

    iface = [["问题", "接口约定"],
             ["何时交接", "需要工程量、综合单价、成本、调价、限价测算时，交造价技能（zaojia-dinge / cost-estimator）"],
             ["交什么", "招标文件清单与计价要求、最高投标限价或控制价、项目规模与工艺方案、拟报价策略与目标得分区间"],
             ["回什么", "分项报价表、综合单价分析表、报价汇总表、报价合理性说明"],
             ["填哪里", "回填《报价表》与报价响应条目，并同步更新评分响应表报价行；报价复核清单逐项勾选留痕"],
             ["不回什么", "不把测算过程文件直接编入投标文件；对外只提交报价表与说明"]]
    return {"摘要": summary, "敏感性": sens_rows, "推演": demo, "复核清单": check_rows,
            "响应条目": resp_rows, "接口": iface}, pts, sens, deducts


def md_of(sec):
    L = ["# 报价评分核对表", "",
         "> 本表由招标文件解析生成，分值、规则、限价均须回原文核对；报价测算须联动造价技能，本表不含测算。", "",
         "## 一、报价评分办法摘要", "", "| " + " | ".join(sec["摘要"][0]) + " |", "|---|---|---|"]
    L += ["| " + " | ".join(str(c) for c in r) + " |" for r in sec["摘要"][1:]]
    L += ["", "## 二、得分敏感性（按规则推演值，非承诺值）", ""]
    if len(sec["敏感性"]) > 1:
        L += ["| " + " | ".join(sec["敏感性"][0]) + " |", "|---|---|---|---|---|"]
        L += ["| " + " | ".join(str(c) for c in r) + " |" for r in sec["敏感性"][1:]]
    else:
        L.append("- 未解析出可计算的扣分规则（口径可能为分档制或公式制），须按招标文件公式手工推演。")
    L += ["", "> 说明：基准价多在开标后依有效报价计算，上表为规则推演，用于报价策略时点判断，非报价测算；"
              "每条规则已标注适用的偏差方向，请勿跨方向套用。", ""]
    L += ["## 三、指定报价得分推演（示例）", ""]
    if sec["推演"]:
        L += ["| " + " | ".join(sec["推演"][0]) + " |", "|---|---|---|---|---|"]
        L += ["| " + " | ".join(str(c) for c in r) + " |" for r in sec["推演"][1:]]
    else:
        L.append("- 未提供 --base/--price，或报价分值未解析，本节省略。")
    L += ["", "## 四、报价复核清单（递交前逐项勾选）", "",
          "| " + " | ".join(sec["复核清单"][0]) + " |", "|---|---|---|---|---|"]
    L += ["| " + " | ".join(str(c) for c in r) + " |" for r in sec["复核清单"][1:]]
    L += ["", "## 五、报价响应条目（五列，可直接并入评分响应表）", "",
          "| " + " | ".join(sec["响应条目"][0]) + " |", "|---|---|---|---|---|"]
    L += ["| " + " | ".join(str(c) for c in r) + " |" for r in sec["响应条目"][1:]]
    L += ["", "> 若使用 skill 脚本生成的响应表骨架（含「原文出处(行号)」辅助列），粘贴时对应前五列、出处列留空即可。", ""]
    L += ["## 六、与造价技能的交接接口", "",
          "| " + " | ".join(sec["接口"][0]) + " |", "|---|---|"]
    L += ["| " + " | ".join(str(c) for c in r) + " |" for r in sec["接口"][1:]]
    L += ["", "> 边界：本技能不出具报价测算结论，不替代造价岗与造价软件成果。"]
    return "\n".join(L)


def dump(sec, out):
    if out.lower().endswith(".csv"):
        import csv as _csv
        with open(out, "w", encoding="utf-8-sig", newline="") as f:
            _csv.writer(f).writerows(sec["响应条目"])
        return
    if out.lower().endswith(".xlsx"):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            print("导出 .xlsx 需 openpyxl：pip install openpyxl", file=sys.stderr)
            sys.exit(3)
        wb = Workbook()
        wb.remove(wb.active)
        fill = PatternFill("solid", fgColor="CCCCCC")
        for name, rows in sec.items():
            ws = wb.create_sheet(name[:31])
            for r in rows:
                ws.append(r)
            for c in ws[1]:
                c.font = Font(bold=True); c.fill = fill
            for i, w in enumerate([22, 56, 24, 30, 22], start=1):
                ws.column_dimensions[chr(64 + i)].width = w
        wb.save(out)
    else:
        open(out, "w", encoding="utf-8").write(md_of(sec))


def main():
    ap = argparse.ArgumentParser(description="报价评分助手（评分规则解析与推演，不含报价测算）")
    ap.add_argument("--file", required=True, help="招标文件文本（txt/md）")
    ap.add_argument("--out", required=True, help="输出（.md / .xlsx / .csv；xlsx 按节分表，csv 输出机读响应条目）")
    ap.add_argument("--base", type=float, help="评标基准价（可选，用于得分推演示例）")
    ap.add_argument("--price", type=float, help="拟投标报价（可选，配合 --base 使用）")
    args = ap.parse_args()

    sec, pts, sens, deducts = build_sections(args.file, args.base, args.price)
    dump(sec, args.out)
    print(f"OK 已生成报价评分核对表：{args.out}")
    print(f"报价分值：{pts:g} 分（L{sec['摘要'][1][2][1:]}）" if pts else "报价分值：未命中，须人工录入")
    print(f"扣分规则解析：{len(deducts)} 条；得分敏感性 {len(sens)} 行（每条规则已标偏差方向）")
    if deducts and not sens:
        print("提示：已解析出扣分规则，但报价分值未命中，敏感性表未生成——请人工录入分值后重算。")
    print("提示：分值/规则/限价均标'待核'，须回原文核对；报价测算须联动造价技能。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
