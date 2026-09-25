#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""招标文件输入预处理 —— 把 PDF / DOC / DOCX / TXT / MD 统一转成脚本可读的 txt/md。

用法:
    python3 tender_input_prep.py --file 招标文件.pdf --out 招标文件.txt
    python3 tender_input_prep.py --file 招标文件.doc --out 招标文件.txt
    python3 tender_input_prep.py --file 招标文件.docx --out 招标文件.md

说明: 图片与扫描件 PDF 本脚本不做 OCR，会提示改用 textin-xparse 等 OCR 技能。
"""
import argparse
import os
import shutil
import subprocess
import sys
import tempfile


def run(cmd, **kw):
    return subprocess.run(cmd, capture_output=True, text=True, **kw)


def docx_extract(path):
    from docx import Document
    d = Document(path)
    parts = [p.text for p in d.paragraphs if p.text and p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def soffice_to_docx(path):
    if shutil.which("soffice") is None:
        print("未检测到 LibreOffice（soffice），无法转换 .doc。"
              "请先安装：apt-get install libreoffice，或用 Word/WPS 另存为 .docx 后重试。", file=sys.stderr)
        sys.exit(3)
    tmpd = tempfile.mkdtemp()
    src = os.path.join(tmpd, "in" + os.path.splitext(path)[1])
    shutil.copy2(path, src)
    run(["soffice", "--headless", "-env:UserInstallation=file:///tmp/lo_prep",
         "--convert-to", "docx", "--outdir", tmpd, src], timeout=180)
    out = os.path.join(tmpd, "in.docx")
    if not os.path.exists(out):
        print("LibreOffice 转换 .doc 失败。请用 Word/WPS 另存为 .docx 后重试。", file=sys.stderr)
        sys.exit(4)
    return out


def page_count(path):
    try:
        import pypdf
        return len(pypdf.PdfReader(path).pages)
    except Exception:
        try:
            import pymupdf as fitz
            return fitz.open(path).page_count
        except Exception:
            return 0


def pdf_extract(path):
    text = ""
    if shutil.which("pdftotext"):
        tmp = tempfile.NamedTemporaryFile(suffix=".txt", delete=False)
        tmp.close()
        run(["pdftotext", "-layout", path, tmp.name])
        if os.path.exists(tmp.name):
            text = open(tmp.name, encoding="utf-8", errors="ignore").read()
            os.remove(tmp.name)
    pages = page_count(path) or 1
    density = len(text.strip()) / max(pages, 1)
    if density < 100:
        try:
            import pymupdf as fitz
            doc = fitz.open(path)
            alt = "\n".join(pg.get_text() for pg in doc)
            if len(alt.strip()) > len(text.strip()):
                text = alt
                density = len(text.strip()) / max(pages, 1)
        except Exception:
            pass
    if density < 100:
        print(f"⚠ 文本密度仅 {density:.0f} 字/页，疑似扫描件或无文字层。"
              "请改用 OCR 技能（如 textin-xparse）转文本后再调用本技能脚本。", file=sys.stderr)
    return text


def main():
    ap = argparse.ArgumentParser(description="招标文件输入预处理")
    ap.add_argument("--file", required=True, help="输入文件（pdf/doc/docx/txt/md）")
    ap.add_argument("--out", required=True, help="输出文件（.txt 或 .md）")
    args = ap.parse_args()

    ext = os.path.splitext(args.file)[1].lower()
    if ext in (".txt", ".md", ".markdown"):
        text = open(args.file, encoding="utf-8", errors="ignore").read()
        kind = "纯文本"
    elif ext == ".docx":
        text = docx_extract(args.file); kind = "docx 直读"
    elif ext == ".doc":
        text = docx_extract(soffice_to_docx(args.file)); kind = "doc 转 docx 后直读"
    elif ext == ".pdf":
        text = pdf_extract(args.file); kind = "pdf 抽文本"
    elif ext in (".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"):
        print("图片输入需 OCR。请使用 textin-xparse 等 OCR 技能转文本后，再调用本技能脚本。", file=sys.stderr)
        sys.exit(3)
    else:
        print(f"不支持的输入格式：{ext}。支持 pdf/doc/docx/txt/md。", file=sys.stderr)
        sys.exit(2)

    text = text.replace("\r\n", "\n").replace("\r", "\n")
    if len(text.strip()) < 50:
        print(f"✗ 提取文本仅 {len(text.strip())} 字，判定为失败（疑似扫描件、加密或缺少 pdftotext/pymupdf）。"
              "请安装依赖或改用 OCR 技能（如 textin-xparse）转文本后重试。未写出输出文件。", file=sys.stderr)
        sys.exit(5)
    with open(args.out, "w", encoding="utf-8") as f:
        f.write(text)
    print(f"OK 已转换：{args.file} → {args.out}（方式：{kind}，{len(text)} 字）")
    print("下一步：python3 bid_red_line_checker.py --file " + args.out + " --out 废标核查表.xlsx")
    return 0


if __name__ == "__main__":
    sys.exit(main())
