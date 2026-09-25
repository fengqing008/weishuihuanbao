#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""编标计划生成器 v1.0 —— 按递交截止日倒排四阶段编标计划（md / xlsx）。

用法:
    python3 bid_plan_builder.py --file 招标文件.md --out 编标计划.xlsx
    python3 bid_plan_builder.py --file 招标文件.md --deadline "2026-10-13 09:00" --out 编标计划.md
"""
import argparse
import csv
import datetime as dt
import re
import sys

D = dt.timedelta(days=1)

# 阶段（偏移天数为相对递交截止日 T）
STAGES = [
    ("阶段一 准备（研读与倒排）", -30, -21, [
        ("获取招标文件并通读研读", "商务/技术负责人", "招标文件研读记录、六要素台账", "递交截止日已锁定"),
        ("判法：确认适用法律与采购方式", "商务负责人", "判法结论（招标投标法/政府采购法）", "采购方式与响应文件口径一致"),
        ("废标红线专项排查", "合规/商务负责人", "废标红线核查表（逐条勾选）", "红线条款全部有证据位置"),
        ("评分办法拆解与得分点倒排", "技术负责人", "得分点清单、资源倒排表", "分值构成与分档已回原文核对"),
        ("资格与业绩资料盘点", "商务负责人", "资格资信材料清单、业绩证明清单", "证照有效期与主体名称一致"),
        ("现场踏勘与答疑问题征集（如组织）", "项目经理", "踏勘记录、答疑问题清单", "踏勘回执已取得"),
        ("异议提出与答疑问题提交", "商务负责人", "书面异议/提问函", "对招标文件有异议的，不迟于投标截止时间 10 日前提出（招标投标法实施条例第二十二条）【提问/答疑截止以招标文件规定为准】"),
    ]),
    ("阶段二 编制（分头成稿）", -20, -11, [
        ("技术标成稿（七板块）", "技术负责人", "技术标正文（含工艺、运营、移交、法律方案）", "重难点措施量化可核查"),
        ("商务标材料归集与撰写", "商务负责人", "资格资信/财务税务社保/业绩/人员/承诺六类材料", "业绩口径与资格条件逐条对应"),
        ("报价编制与复核", "造价岗（联动造价技能）", "报价表、报价合理性说明", "不超最高投标限价、大小写一致"),
        ("评分响应表填写与覆盖度校验", "技术/商务负责人", "五列评分响应表", "覆盖度校验 PASS（未覆盖=0）"),
        ("投标文件初稿汇总与目录编排", "编标专员", "初稿合册", "章节与本项目一致，无串项目内容"),
    ]),
    ("阶段三 审核（内部三级审核）", -10, -2, [
        ("技术审核：方案完整性、偏离表、参数一致", "技术负责人", "技术审核记录", "无负偏离未声明"),
        ("商务审核：主体信息、证照、业绩、报价", "商务负责人", "商务审核记录", "主体名称/证照号全篇一致"),
        ("合规审核：废标红线逐条复核", "合规岗", "红线复核表", "任一不满足项为零"),
        ("签署用印：法定代表人/授权代理人签章", "综合岗", "签章完成的投标文件", "授权委托书与签字人一致"),
        ("装订、密封、标记（正副本/骑缝章）", "编标专员", "密封完成的投标文件", "密封与标记符合须知要求"),
        ("电子投标：CA 锁加密上传与验签（如适用）", "编标专员", "电子标上传成功回执", "不加密与加密版均可用"),
        ("投标保证金提交（转账/保函）", "财务岗", "付款凭证或电子保函", "到账时间不迟于招标文件规定【待核：以招标文件为准】"),
    ]),
    ("阶段四 递交与开标", -1, 0, [
        ("递交前终检（正副本份数、原件、电子件）", "编标专员", "递交清单核对表", "份数与介质符合要求"),
        ("递交投标文件（现场或电子平台）", "授权代理人", "递交回执/上传回执", "在递交截止时间前完成"),
        ("参加开标会、身份核验与签到", "授权代理人", "开标记录、唱标价记录", "核验材料齐全"),
        ("澄清与答疑应对（如被要求）", "技术/商务负责人", "书面澄清回复", "不改变投标文件实质性内容"),
        ("结果跟进与保证金退还跟踪", "商务负责人", "中标/结果通知；退款跟踪表", "工程招标自合同签订后 5 日内退还（招标投标法实施条例第五十七条）"),
    ]),
]


def parse_deadline(text):
    """从招标文件提取递交截止/开标时间。"""
    pats = [
        r"投标截止[^\n。；]{0,10}?(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日\s*(\d{1,2})[:：](\d{2})?",
        r"递交[^\n。；]{0,12}截止[^\n。；]{0,12}?(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日\s*(\d{1,2})[:：](\d{2})?",
        r"开标[^\n。；]{0,12}?(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日\s*(\d{1,2})[:：](\d{2})?",
    ]
    for p in pats:
        m = re.search(p, text)
        if m:
            y, mo, d, h, mi = m.groups()
            return dt.datetime(int(y), int(mo), int(d), int(h), int(mi or 0)), m.group(0).strip()
    return None, ""


def parse_date_arg(s):
    for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d", "%Y/%m/%d %H:%M", "%Y/%m/%d"):
        try:
            return dt.datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def build_rows(deadline):
    rows = [["序号", "阶段", "任务", "责任岗位", "开始日期", "完成日期", "交付物", "检查点"]]
    n = 1
    for stage, off_s, off_e, tasks in STAGES:
        s_date, e_date = deadline + off_s * D, deadline + off_e * D
        span = len(tasks)
        for k, (task, owner, deliver, chk) in enumerate(tasks):
            d0 = s_date + dt.timedelta(days=(e_date - s_date).days * k // span)
            d1 = s_date + dt.timedelta(days=(e_date - s_date).days * (k + 1) // span)
            rows.append([n, stage, task, owner, d0.strftime("%Y-%m-%d"), d1.strftime("%Y-%m-%d"), deliver, chk])
            n += 1
    return rows


def write_md(rows, out, deadline, src_note):
    L = ["# 编标计划（按递交截止日倒排）", "",
         f"> 递交截止（T）：{deadline.strftime('%Y-%m-%d %H:%M')}　来源：{src_note}", "",
         f"> 关键词点：澄清/修改影响投标文件编制的，应不迟于 T-15 日发出（招标投标法实施条例第二十一条、政府采购法实施条例第三十一条）；"
         f"对招标文件有异议的，应不迟于 T-10 日提出（招标投标法实施条例第二十二条）【提问/答疑截止以招标文件规定为准】", "",
         "| " + " | ".join(rows[0]) + " |", "|" + "---|" * len(rows[0])]
    for r in rows[1:]:
        L.append("| " + " | ".join(str(c) for c in r) + " |")
    L += ["", "> 责任岗位与日期为倒排建议，须按投标单位实际人员与工作安排调整后执行。"]
    open(out, "w", encoding="utf-8").write("\n".join(L))


def write_csv(rows, out):
    with open(out, "w", encoding="utf-8-sig", newline="") as f:
        csv.writer(f).writerows(rows)


def write_xlsx(rows, out):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        print("导出 .xlsx 需 openpyxl：pip install openpyxl", file=sys.stderr)
        sys.exit(3)
    wb = Workbook(); ws = wb.active; ws.title = "编标计划"
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="CCCCCC")
        c.alignment = Alignment(horizontal="center", vertical="center")
    for col, w in zip("ABCDEFGH", [5, 22, 34, 18, 12, 12, 34, 40]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb.save(out)


def main():
    ap = argparse.ArgumentParser(description="编标计划生成器")
    ap.add_argument("--file", help="招标文件文本（txt/md），用于自动抓取递交截止时间")
    ap.add_argument("--deadline", help='递交截止时间，如 "2026-10-13 09:00"（提供则覆盖文件抓取值）')
    ap.add_argument("--out", required=True, help="输出计划（.md / .csv / .xlsx）")
    args = ap.parse_args()

    note = ""
    deadline = parse_date_arg(args.deadline) if args.deadline else None
    if deadline:
        note = "命令参数指定"
    elif args.file:
        try:
            with open(args.file, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
        except OSError as e:
            print(f"读取招标文件失败：{e}", file=sys.stderr)
            return 2
        deadline, raw = parse_deadline(text)
        note = f"招标文件原文「{raw}」" if deadline else ""
    if deadline is None:
        print("未能确定递交截止时间。请用 --deadline \"YYYY-MM-DD HH:MM\" 指定后重试。"
              "（自动抓取失败常见于扫描件或前附表式日期格式，属正常，需人工补录）", file=sys.stderr)
        return 2

    print("== 编标计划关键日期（倒排）==")
    print(f"递交截止/开标（T）：{deadline.strftime('%Y-%m-%d %H:%M')}　[{note}]")
    print(f"澄清/修改发出最后日（T-15）：{deadline - 15 * D:%Y-%m-%d}（招标投标法实施条例第二十一条）")
    print(f"异议提出最后日（T-10）：{deadline - 10 * D:%Y-%m-%d}（招标投标法实施条例第二十二条；提问/答疑截止以招标文件规定为准）")
    print(f"编制启动建议（T-20）：{deadline - 20 * D:%Y-%m-%d}")
    print(f"内部审核启动建议（T-10）：{deadline - 10 * D:%Y-%m-%d}")

    rows = build_rows(deadline)
    low = args.out.lower()
    if low.endswith(".xlsx"):
        write_xlsx(rows, args.out)
    elif low.endswith(".csv"):
        write_csv(rows, args.out)
    else:
        write_md(rows, args.out, deadline, note)
    print(f"OK 编标计划已写入 {args.out}（四阶段 {len(rows)-1} 项任务，含责任岗位/日期/交付物/检查点）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
