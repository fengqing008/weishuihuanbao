#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""评分办法拆解器 v3 —— 行式 + 前附表式双通道 + 原文出处 + 分值合计。

用法:
    python3 scoring_matrix.py --file 招标文件.md --out 得分点清单.xlsx
"""
import argparse
import csv
import re
import sys

BUCKETS = {
    "技术分": ["技术分", "技术部分", "技术方案", "技术评审", "技术标", "技术能力",
               "运行维护", "重难点", "质量保证", "应急处理", "自检能力", "服务方案", "运营方案"],
    "商务分": ["商务分", "商务部分", "综合实力", "商务评审", "资信", "业绩", "企业实力",
               "项目团队", "服务承诺"],
    "财务融资分": ["财务分", "财务融资", "融资方案", "财务状况", "财务方案"],
    "法律分": ["法律分", "法律方案", "法律服务"],
    "报价分": ["报价分", "价格分", "报价部分", "投标报价", "价格评审"],
}

TEMPLATE_TECH = [
    ("项目理解与总体筹划", "项目概况、背景、服务范围、总体目标、优化建议"),
    ("项目公司组建与管理", "组建方案、组织架构、管理职能、人员配置、资本金与融资安排"),
    ("设计方案（工艺）", "工艺比选、设计参数、去除率、设备清单"),
    ("建设方案（施工组织）", "施工部署、工期安排、质量安全、成本控制"),
    ("运营方案", "组织定员、水质保障、设备管网运维、成本考核、应急"),
    ("移交方案", "移交范围、恢复性大修、人员培训、缺陷责任期"),
    ("法律方案", "协议偏差表、条款响应"),
]
TEMPLATE_BIZ = [
    ("投标核心文件", "投标函、承诺书、报价表、报价合理性分析"),
    ("主体资格与授权", "营业执照、法定代表人身份证明、授权书、资格承诺书"),
    ("财务税务社保", "财务状况报告、审计报告、纳税证明、社保证明"),
    ("业绩", "近5年类似项目业绩证明、业绩表"),
    ("人员与机构", "项目负责人证书、团队名单、劳动合同、资格证书"),
    ("声明与承诺", "无重大违法记录、无失信记录、无拖欠工资、廉洁、保密"),
]


def cell_val(c):
    """单元格取值：支持 '30' 与 '30分' 两种写法；非纯数值返回 None。"""
    m = re.fullmatch(r"(\d{1,3}(?:\.\d{1,2})?)\s*分?", c or "")
    return float(m.group(1)) if m else None


def fmt(v):
    return int(v) if float(v) == int(v) else v


def scan(path):
    """双通道抓取：①前附表/评审因素表（markdown 表格行）②正文行式（关键词+数字+分）。"""
    hits = {k: [] for k in BUCKETS}
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        for i, line in enumerate(f, 1):
            s = line.strip()
            if s.startswith("|"):
                cells = [c.strip() for c in s.strip("|").split("|")]
                if all(set(c) <= set("-: ") for c in cells):
                    continue
                vals = [v for v in (cell_val(c) for c in cells) if v is not None]
                if vals and "分" in s:
                    cand = vals[1:] if len(vals) > 1 else vals
                    val = cand[0]
                    if 0 < val <= 100:
                        for bucket, kws in BUCKETS.items():
                            if any(kw in s for kw in kws):
                                hits[bucket].append((fmt(val), i, s[:140]))
                                break
                        continue
            for bucket, kws in BUCKETS.items():
                for kw in kws:
                    m = re.search(re.escape(kw) + r"[^\n。；;|]{0,20}?(\d{1,3})\s*分", line)
                    if m:
                        hits[bucket].append((int(m.group(1)), i, line.strip()[:140]))
                        break
    return hits


def sums(hits):
    return {k: sum(v for v, _, _ in items) for k, items in hits.items()}


def build_rows(hits):
    total = sums(hits)
    rows = [["类别", "评审因素/分值", "分值(机抓·待核)", "原文出处(行号)", "得分点要点", "责任章节"]]
    print("== 分值构成（机抓，须回原文核对分档与例外）==")
    grand = 0
    for bucket in ["技术分", "商务分", "财务融资分", "法律分", "报价分"]:
        items = hits.get(bucket, [])
        if items:
            vals = "/".join(str(v) for v, _, _ in items)
            lines = ",".join(f"L{l}" for _, l, _ in items)
            grand += total[bucket]
            print(f"{bucket}: {vals}　合计 {total[bucket]}　出处 {lines}")
        else:
            vals, lines = "待核", ""
            print(f"{bucket}: 未命中（须人工录入）")
        rows.append([bucket, f"{bucket}（机抓分项）", vals, lines, "", ""])
    print(f"机抓分值合计: {grand}（与招标文件标注满分核对，不一致处以招标文件为准）")
    rows.append(["合计", "机抓分值合计", grand, "", "须与招标文件满分核对", ""])
    rows.append([])
    rows.append(["【得分点准备清单·技术标（七板块）】", "", "", "", "", ""])
    for name, point in TEMPLATE_TECH:
        rows.append(["技术标", name, "待核", "", point, name])
    rows.append([])
    rows.append(["【得分点准备清单·商务标】", "", "", "", "", ""])
    for name, point in TEMPLATE_BIZ:
        rows.append(["商务标", name, "待核", "", point, name])
    return rows


def write_out(rows, out):
    if out.lower().endswith(".xlsx"):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            print("导出 .xlsx 需 openpyxl：pip install openpyxl", file=sys.stderr)
            sys.exit(3)
        wb = Workbook(); ws = wb.active; ws.title = "得分点清单"
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="CCCCCC")
        for col, w in zip("ABCDEF", [14, 26, 18, 20, 46, 20]):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"
        wb.save(out)
    else:
        with open(out, "w", encoding="utf-8-sig", newline="") as f:
            csv.writer(f).writerows(rows)


def main():
    ap = argparse.ArgumentParser(description="评分办法拆解器 v3（行式+前附表式）")
    ap.add_argument("--file", required=True, help="招标文件文本（txt/md）")
    ap.add_argument("--out", required=True, help="输出清单（.csv / .xlsx）")
    args = ap.parse_args()
    hits = scan(args.file)
    write_out(build_rows(hits), args.out)
    print(f"OK 清单已写入 {args.out}（分值列为机抓值，附原文行号，须回招标文件核对）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
