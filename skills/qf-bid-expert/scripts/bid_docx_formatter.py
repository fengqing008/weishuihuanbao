#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""标书 Word 排版器 —— Markdown/Word 转为规范投标文件：
标题自动挂接 Word 大纲级别（Heading 1/2/3）+ 插入可一键更新的 TOC 目录域。

用法:
    python3 bid_docx_formatter.py -i 技术标.md -o 技术标.docx --type technical
    python3 bid_docx_formatter.py -i 商务标.md -o 商务标.docx --type commercial
    python3 bid_docx_formatter.py --verify 技术标.docx
"""
import argparse
import os
import re
import subprocess
import sys
import tempfile

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
    from docx.text.paragraph import Paragraph
except ImportError:
    print("缺少依赖 python-docx。请先安装（在本技能目录下执行）：\n"
          "    pip install -r requirements.txt\n"
          "或：pip install python-docx", file=sys.stderr)
    sys.exit(3)

# type -> (正文字体, 标题字体, 正文字号pt, 一级标题字号pt, 行距pt, 页边距cm)
PROFILES = {
    "technical": ("仿宋", "仿宋", 12, 16, 25, (2.5, 2.0, 2.0, 2.0)),
    "commercial": ("宋体", "黑体", 12, 16, 22, (2.54, 2.54, 3.17, 3.17)),
}
EN_FONT = "Times New Roman"


# ----------------------------------------------------------- 基础工具
def _style(doc, name):
    for s in doc.styles:
        if (s.name or "").lower() == name.lower():
            return s
    return None


def set_style_font(style, cn, size=None, bold=None):
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    style.font.color.rgb = None
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:ascii"), EN_FONT)
    rf.set(qn("w:hAnsi"), EN_FONT)
    rf.set(qn("w:eastAsia"), cn)


def set_style_para(style, before=None, after=None, line_pt=None, align=None, indent_chars=None):
    p = style.paragraph_format
    if before is not None:
        p.space_before = Pt(before)
    if after is not None:
        p.space_after = Pt(after)
    if line_pt is not None:
        p.line_spacing = Pt(line_pt)
    if align is not None:
        p.alignment = align
    if indent_chars is not None:
        pPr = style.element.get_or_add_pPr()
        ind = pPr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            pPr.append(ind)
        ind.set(qn("w:firstLineChars"), str(indent_chars * 100))


def add_toc_field(paragraph, levels="1-3"):
    run = paragraph.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); run._r.append(b)
    t = OxmlElement("w:instrText"); t.set(qn("xml:space"), "preserve")
    t.text = 'TOC \\o "%s" \\h \\z \\u' % levels
    run._r.append(t)
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate"); run._r.append(s)
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); run._r.append(e)


def set_update_fields(doc):
    try:
        settings = doc.settings.element
    except Exception:
        return
    if settings.find(qn("w:updateFields")) is None:
        el = OxmlElement("w:updateFields"); el.set(qn("w:val"), "true")
        settings.append(el)


def style_page(doc, margins):
    top, bottom, left, right = margins
    for sec in doc.sections:
        sec.top_margin = Cm(top); sec.bottom_margin = Cm(bottom)
        sec.left_margin = Cm(left); sec.right_margin = Cm(right)


def style_footer(section):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, typ in [("w:fldChar", "begin"), ("w:instrText", None), ("w:fldChar", "end")]:
        el = OxmlElement(tag)
        if typ:
            el.set(qn("w:fldCharType"), typ)
        else:
            el.set(qn("xml:space"), "preserve"); el.text = "PAGE"
        run._r.append(el)


def shift_headings(doc):
    mapping = {"Heading 1": "Title", "Heading 2": "Heading 1",
               "Heading 3": "Heading 2", "Heading 4": "Heading 3"}
    for p in list(doc.paragraphs):
        sn = (p.style.name or "")
        if sn in mapping:
            st = _style(doc, mapping[sn])
            if st is not None:
                p.style = st


def remove_existing_toc(doc):
    for p in list(doc.paragraphs):
        if p.text.strip().replace("　", "") == "目录" and (p.style.name or "").startswith(("Heading", "Title")):
            p._element.getparent().remove(p._element)


def style_styles(doc, prof):
    cn, hcn, body_sz, h1_sz, line_pt, _ = PROFILES[prof]
    t = _style(doc, "Title")
    if t:
        set_style_font(t, hcn, h1_sz + 4, True)
        set_style_para(t, before=0, after=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    for name, size, align, indent in [("Heading 1", h1_sz, WD_ALIGN_PARAGRAPH.CENTER, 0),
                                      ("Heading 2", body_sz, WD_ALIGN_PARAGRAPH.LEFT, 2),
                                      ("Heading 3", body_sz, WD_ALIGN_PARAGRAPH.LEFT, 2)]:
        st = _style(doc, name)
        if st:
            set_style_font(st, hcn, size, True)
            set_style_para(st, before=6, after=6, line_pt=line_pt, align=align, indent_chars=indent)
    for name in ["Normal", "Body Text"]:
        st = _style(doc, name)
        if st:
            set_style_font(st, cn, body_sz, False)
            set_style_para(st, line_pt=line_pt, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent_chars=2)


def style_tables(doc):
    for t in doc.tables:
        tblPr = t._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement("w:" + edge)
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "6")
            el.set(qn("w:color"), "000000")
            borders.append(el)
        tblPr.append(borders)


def convert_md(path):
    import shutil
    if shutil.which("pandoc") is None:
        print("未检测到 pandoc（Markdown → docx 转换需要）。\n"
              "请先安装：apt-get install pandoc（或 brew/choco install pandoc），\n"
              "或将底稿另存为 .docx 后以 -i 文件名.docx 方式调用。", file=sys.stderr)
        sys.exit(3)
    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    tmp.close()
    cmd = ["pandoc", path, "-o", tmp.name, "-f", "markdown-smart"]
    r = subprocess.run(cmd, capture_output=True, text=True)
    if r.returncode != 0:
        raise RuntimeError("pandoc 转换失败: " + (r.stderr or "")[:200])
    return tmp.name


def verify(path):
    doc = Document(path)
    h = {"Heading 1": 0, "Heading 2": 0, "Heading 3": 0}
    for p in doc.paragraphs:
        sn = p.style.name or ""
        if sn in h:
            h[sn] += 1
    has_toc = "TOC" in doc.element.xml.upper()
    try:
        has_upd = doc.settings.element.find(qn("w:updateFields")) is not None
    except Exception:
        has_upd = False
    print("== 大纲与目录校验 ==")
    print(f"Heading 1 段落数: {h['Heading 1']}")
    print(f"Heading 2 段落数: {h['Heading 2']}")
    print(f"Heading 3 段落数: {h['Heading 3']}")
    print(f"含 TOC 域: {has_toc}")
    print(f"updateFields 已设: {has_upd}")
    ok = (h["Heading 1"] + h["Heading 2"] + h["Heading 3"] > 0) and has_toc and has_upd
    print("结论:", "PASS（可一键生成目录）" if ok else "FAIL（检查标题样式或 TOC 域）")
    return 0 if ok else 1


def main():
    ap = argparse.ArgumentParser(description="标书 Word 排版器")
    ap.add_argument("-i", "--input", help="输入 md 或 docx")
    ap.add_argument("-o", "--output", help="输出 docx")
    ap.add_argument("--type", choices=["technical", "commercial"], default="technical", help="标书类型")
    ap.add_argument("--no-toc", action="store_true", help="不插目录域")
    ap.add_argument("--no-shift", action="store_true", help="不执行层级下移")
    ap.add_argument("--verify", help="校验既有 docx 的大纲级别与目录域")
    args = ap.parse_args()

    if args.verify:
        return verify(args.verify)
    if not args.input or not args.output:
        print("需提供 -i 与 -o（或使用 --verify）", file=sys.stderr)
        return 2

    src = args.input
    tmp = None
    if src.lower().endswith((".md", ".markdown", ".txt")):
        tmp = convert_md(src)
        src = tmp
    doc = Document(src)
    if not args.no_shift:
        shift_headings(doc)
    remove_existing_toc(doc)
    style_page(doc, PROFILES[args.type][5])
    style_styles(doc, args.type)
    style_tables(doc)
    for sec in doc.sections:
        style_footer(sec)
    if not args.no_toc:
        anchor = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()

        def add_after(a):
            np = OxmlElement("w:p")
            a._element.addnext(np)
            return Paragraph(np, a._parent)

        head = add_after(anchor)
        head.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hr = head.add_run("目　录")
        hr.bold = True
        hr.font.size = Pt(PROFILES[args.type][3])
        field_p = add_after(head)
        add_toc_field(field_p)
        set_update_fields(doc)
    doc.save(args.output)
    if tmp and os.path.exists(tmp):
        os.remove(tmp)
    print(f"OK 已生成 {args.output}（{args.type} 档，标题挂 Heading 1/2/3，含一键目录域）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
