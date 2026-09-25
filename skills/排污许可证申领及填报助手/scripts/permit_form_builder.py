#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
permit_form_builder.py — 排污许可证申领及填报助手核心脚本

四个子命令：
  classify       管理类别判定（重点 / 简化 / 登记），附依据条文
  checklist      排污许可申报资料清单（Markdown / CSV / XLSX）
  limit-calc     许可排放量核算（COD / 氨氮 / 总氮 / 总磷 多指标，含取严）
  form-template  申请表填报模板（字段以【待填：xxx】占位）

依赖：Python 3.8+ 标准库。openpyxl 为可选增强（仅 --format xlsx 时使用），
缺失时自动降级为 CSV 并给出提示，不中断执行。

依据：
  《排污许可管理条例》（国务院令第736号）
  《排污许可管理办法》（生态环境部令第32号，2024-07-01 施行）
  《固定污染源排污许可分类管理名录（2019年版）》（生态环境部令第11号）
  HJ 978-2018《排污许可证申请与核发技术规范 水处理（试行）》
  HJ 944-2018《排污单位环境管理台账及排污许可证执行报告技术规范 总则（试行）》

用法示例：
  python3 permit_form_builder.py --help
  python3 permit_form_builder.py classify --industry 污水处理及其再生利用 --scale 5000 --water-type domestic --discharge direct
  python3 permit_form_builder.py checklist --management key --format md --out ./permit_checklist.md
  python3 permit_form_builder.py limit-calc --design-flow 5000 --hours 8760 --pollutants COD,氨氮,总氮,总磷 --limits 50,8,15,0.5
  python3 permit_form_builder.py form-template --type wastewater --management key --out ./permit_form.md
"""

import argparse
import csv
import io
import os
import sys

VERSION = "1.0.0"

# ── 行业与类别常量 ────────────────────────────────────────────────
INDUSTRY_ALIAS = {
    "污水处理及其再生利用": "462",
    "污水处理": "462",
    "城镇污水处理厂": "462",
    "工业废水集中处理厂": "462",
    "水处理通用工序": "112",
    "462": "462",
    "112": "112",
}

CATEGORY_KEY = "重点管理"
CATEGORY_SIMPLE = "简化管理"
CATEGORY_REG = "登记管理"

# 城镇污水处理厂污染物排放标准 GB 18918-2002 一级A 常用限值（mg/L）
GB18918_1A = {
    "COD": 50.0,
    "化学需氧量": 50.0,
    "BOD5": 10.0,
    "SS": 10.0,
    "悬浮物": 10.0,
    "氨氮": 5.0,
    "总氮": 15.0,
    "总磷": 0.5,
}

DISCHARGE_TEXT = {
    "direct": "直接排入环境水体",
    "indirect": "间接排放（排入城镇污水收集系统或受纳单位）",
    "reuse": "出水再生利用，不外排",
}

WATER_TYPE_TEXT = {
    "domestic": "城乡污水集中处理场所（生活污水为主）",
    "industrial": "工业废水集中处理场所",
    "mixed": "混合废水（生活污水与工业废水兼收）",
}


# ── 工具函数 ──────────────────────────────────────────────────────
def _split_list(text):
    """把 'a,b,c' 拆成列表，兼容中英文逗号与顿号。"""
    if not text:
        return []
    for sep in ("，", "、", ";"):
        text = text.replace(sep, ",")
    return [x.strip() for x in text.split(",") if x.strip()]


def _emit(content, out_path=None):
    """把内容输出到 stdout；给定路径时同时落盘（自动建目录）。"""
    if out_path:
        out_dir = os.path.dirname(os.path.abspath(out_path))
        if out_dir and not os.path.isdir(out_dir):
            os.makedirs(out_dir, exist_ok=True)
        with open(out_path, "w", encoding="utf-8") as fh:
            fh.write(content)
        print(content)
        print("\n[OK] 已写入：%s（%d 字）" % (out_path, len(content)))
    else:
        print(content)


def _norm_pollutant(name):
    """污染物别名归一到标准写法。"""
    table = {
        "cod": "化学需氧量（CODcr）",
        "codcr": "化学需氧量（CODcr）",
        "化学需氧量": "化学需氧量（CODcr）",
        "氨氮": "氨氮（以N计）",
        "nh3n": "氨氮（以N计）",
        "总氮": "总氮（以N计）",
        "tn": "总氮（以N计）",
        "总磷": "总磷（以P计）",
        "tp": "总磷（以P计）",
        "ph": "pH值",
        "ss": "悬浮物",
        "悬浮物": "悬浮物",
        "bod5": "五日生化需氧量（BOD5）",
        "五日生化需氧量": "五日生化需氧量（BOD5）",
        "色度": "色度",
        "石油类": "石油类",
        "动植物油": "动植物油",
    }
    key = name.strip().lower()
    if key in table:
        return table[key]
    return name.strip()


# GB 18918-2002 一级A 默认限值，键为归一后的标准写法（mg/L）
DEFAULT_LIMIT = {
    "化学需氧量（CODcr）": 50.0,
    "氨氮（以N计）": 8.0,   # 限值 5(8)：水温大于12℃按5，小于等于12℃按8，核算年许可量按8
    "总氮（以N计）": 15.0,
    "总磷（以P计）": 0.5,
    "悬浮物": 10.0,
    "五日生化需氧量（BOD5）": 10.0,
    "色度": 30.0,
    "石油类": 1.0,
    "动植物油": 1.0,
}


def _default_limit(pollutant):
    """按归一名称取 GB 18918-2002 一级A 默认浓度限值。"""
    return DEFAULT_LIMIT.get(_norm_pollutant(pollutant), 0.0)


# ── 子命令 1：classify ────────────────────────────────────────────
def cmd_classify(args):
    industry_code = INDUSTRY_ALIAS.get(args.industry.strip(), None)
    scale = args.scale
    wt = args.water_type
    lines = []
    lines.append("# 排污许可管理类别判定意见")
    lines.append("")
    lines.append("| 项目 | 输入值 |")
    lines.append("|---|---|")
    lines.append("| 行业类别（申报口径） | %s |" % args.industry)
    lines.append("| 行业代码 | %s |" % (industry_code or "未识别，须人工核对《国民经济行业分类》GB/T 4754"))
    lines.append("| 设计处理能力 | %.0f m³/d |" % scale)
    lines.append("| 污水来源类型 | %s |" % WATER_TYPE_TEXT.get(wt, wt))
    lines.append("| 排放去向 | %s |" % DISCHARGE_TEXT.get(args.discharge, args.discharge))
    lines.append("")

    category = None
    basis = []
    actions = []

    if industry_code == "112":
        category = CATEGORY_KEY
        basis.append(
            "《固定污染源排污许可分类管理名录（2019年版）》第112项「水处理通用工序」"
            "为兜底类别：第1至108类行业中明确涉及水处理通用工序的排污单位，"
            "其水处理工段按本项判定，并适用《排污许可证申请与核发技术规范 水处理通用工序》。"
        )
        actions.append("按水处理通用工序技术规范单独核算水处理工段的许可要求；主行业要求另按主行业规范执行。")
    elif industry_code == "462":
        if wt == "industrial":
            category = CATEGORY_KEY
            basis.append(
                "《固定污染源排污许可分类管理名录（2019年版）》第99项「污水处理及其再生利用（462）」："
                "「工业废水集中处理场所」列入**重点管理**，与日处理能力规模无关。"
            )
            actions.append("按重点管理申领排污许可证；申领前须在全国排污许可证管理信息平台公开基本信息与拟申请许可事项不少于5个工作日。")
            actions.append("许可排放量须明确化学需氧量、氨氮、总氮、总磷四项。")
        else:
            if scale >= 20000:
                category = CATEGORY_KEY
                basis.append(
                    "《固定污染源排污许可分类管理名录（2019年版）》第99项："
                    "「日处理能力2万吨及以上的城乡污水集中处理场所」列入**重点管理**。"
                )
                actions.append("按重点管理申领排污许可证；须安装自动监测设备并与生态环境主管部门监控系统联网。")
            elif scale >= 500:
                category = CATEGORY_SIMPLE
                basis.append(
                    "《固定污染源排污许可分类管理名录（2019年版）》第99项："
                    "「日处理能力500吨及以上、2万吨以下的城乡污水集中处理场所」列入**简化管理**。"
                )
                actions.append("按简化管理申领排污许可证；提交年度执行报告与季度执行报告，台账内容可按 HJ 978-2018 适当缩减。")
            else:
                category = CATEGORY_REG
                basis.append(
                    "《固定污染源排污许可分类管理名录（2019年版）》第99项："
                    "「日处理能力500吨以下」列入**登记管理**，填报排污登记表，不属行政许可。"
                )
                actions.append("在全国排污许可证管理信息平台填报排污登记表，提交后即时生成登记编号与回执；信息变动自变动之日起20日内变更登记。")
            if wt == "mixed":
                basis.append(
                    "兼收工业废水的城乡污水集中处理场所，仍按城乡污水集中处理场所的规模档判定；"
                    "若受纳园区工业废水且性质为工业废水集中处理场所，则改按工业废水集中处理场所执行重点管理。"
                )
    else:
        lines.append("🔴 未能识别行业类别，停止自动判定。请核对是否为「污水处理及其再生利用（462）」或「水处理通用工序（112）」。")
        lines.append("")
        lines.append("请改用以下命令重跑：")
        lines.append("")
        lines.append("```")
        lines.append(
            "python3 scripts/permit_form_builder.py classify --industry 污水处理及其再生利用 "
            "--scale 20000 --water-type industrial --discharge direct"
        )
        lines.append("```")
        _emit("\n".join(lines), args.out)
        return 2

    # 规模以设计处理能力为准的补充口径
    basis.append(
        "规模口径：以排污单位实际设计处理能力为准；实际设计处理能力与环评批复设计处理能力不一致时，以实际设计处理能力为准。"
    )
    if args.cod_total is not None and category == CATEGORY_REG:
        basis.append(
            "_名录_第七条「化学需氧量年排放量大于30吨」等按排放量纳管的条款仅适用于第108项「其他行业」，"
            "污水处理及其再生利用属第99项，不适用该条，故本判定不因 COD 年排放量调档。"
        )

    # 再生利用对许可内容的影响
    limit_note = []
    if args.discharge == "reuse":
        if wt == "industrial":
            limit_note.append("工业废水集中处理厂出水再生利用：许可排放浓度与许可排放量均不许可，须在副本中说明去向与受纳单位。")
        else:
            limit_note.append("城镇及其他生活污水处理厂出水再生利用：仅许可排放浓度，不许可排放量。")
    else:
        limit_note.append("出水直接或间接排入环境水体：废水排放口全部为主要排放口，许可排放浓度与年许可排放量均须明确。")

    lines.append("## 判定结论")
    lines.append("")
    lines.append("| 项目 | 结论 |")
    lines.append("|---|---|")
    lines.append("| 管理类别 | **%s** |" % category)
    lines.append("| 管理形式 | %s |" % (
        "申请取得排污许可证（行政许可）" if category != CATEGORY_REG else "填报排污登记表（非行政许可）"
    ))
    lines.append("| 有效期 | %s |" % ("5年，届满前60日申请延续" if category != CATEGORY_REG else "随登记信息变动20日内变更，不设有效期"))
    lines.append("| 执行报告 | %s |" % (
        "年度执行报告 + 季度执行报告" if category == CATEGORY_KEY
        else ("年度执行报告 + 季度执行报告（简化内容）" if category == CATEGORY_SIMPLE else "不适用")
    ))
    lines.append("")
    lines.append("## 判定依据")
    lines.append("")
    for i, b in enumerate(basis, 1):
        lines.append("%d. %s" % (i, b))
    lines.append("")
    lines.append("## 后续动作清单")
    lines.append("")
    for i, a in enumerate(actions, 1):
        lines.append("%d. %s" % (i, a))
    lines.append("")
    lines.append("## 许可排放限值口径")
    lines.append("")
    for i, n in enumerate(limit_note, 1):
        lines.append("%d. %s" % (i, n))
    lines.append("")
    lines.append("🔴 本次判定为脚本按名录规则自动给出，须在提交申报前与地方生态环境主管部门的核发口径核对一次，")
    lines.append("并以核发机关书面答复为准。")
    _emit("\n".join(lines), args.out)
    return 0


# ── 子命令 2：checklist ───────────────────────────────────────────
CHECKLIST_COMMON = [
    ("1", "排污许可证申请表（平台生成并导出 PDF）", "必备", "1份", "法定代表人签字并加盖公章；平台填报内容与纸质件一致"),
    ("2", "营业执照或事业单位法人证书副本", "必备", "1份", "统一社会信用代码须与平台账号一致，复印件加盖公章"),
    ("3", "法定代表人身份证正反面复印件", "必备", "1份", "委托办理时另附授权委托书与经办人身份证"),
    ("4", "建设项目环境影响报告书（表）批复文件", "必备", "1份", "或环境影响登记表备案回执；2015-01-01（含）后取得批复的，许可排放量须同时满足批复要求"),
    ("5", "承诺书（完整真实合法、按证排污）", "必备", "1份", "法定代表人或主要负责人签字或盖章；基本信息变更与延续情形可不提交"),
    ("6", "排放口规范化情况说明", "必备", "1份", "已建成排放口须提交；含排放口编号、坐标、标志牌与采样口设置情况"),
    ("7", "自行监测方案", "必备", "1份", "含点位及示意图、指标、频次、分析方法、质量保证与质量控制、数据记录与信息公开放"),
    ("8", "许可排放量限值计算过程", "必备", "1份", "申请许可排放量的一并提交；含计算式、参数来源与取严过程"),
    ("9", "污水处理设施纳污范围说明", "必备", "1份", "城镇污水集中处理设施与工业废水集中处理设施均须提交；含管网布置与最终排放去向"),
    ("10", "废水排放口地理坐标与受纳水体信息", "必备", "1份", "直接排放填报入河排污口名称编号、受纳水体功能目标与汇入处坐标"),
    ("11", "厂区总平面布置图与污水处理工艺流程图", "必备", "1份", "标注产排污环节编号（MF/TW/TS/TA），与平台填报编号一致"),
    ("12", "主要生产设施与污染防治设施技术参数表", "必备", "1份", "含设计值、设计处理能力、年运行小时数"),
    ("13", "药剂投加量与投加点位说明", "必备", "1份", "PAC、PAM、次氯酸钠、双氧水、硫酸、氢氧化钠等，按工艺单元列设计投加量"),
    ("14", "污泥处理处置去向与委托合同", "必备", "1份", "含处理后含水率、厂内暂存量、综合利用量、委托处置单位与资质"),
    ("15", "固体废物分类与去向清单", "必备", "1份", "含一般工业固废（污泥 SW07）与危险废物（HW49 在线监测废液、HW08 废润滑油）"),
    ("16", "重点污染物总量控制指标来源说明", "必备（重点管理）", "1份", "含通过削减替代获得指标的说明材料"),
    ("17", "排污权交易指标证明材料", "按需", "1份", "重点污染物总量指标通过排污权交易获取时提交"),
    ("18", "平台信息公开情况说明", "必备（重点管理）", "1份", "首次申请或重新申请前公开不少于5个工作日，附公开截图与起止时间"),
    ("19", "近三年实际排水量与运行小时数统计表", "必备", "1份", "用于确定年许可排放量的水量基数；运行不满3年从投产之日起算"),
    ("20", "进水水质监测数据与进水自动监测联网情况说明", "必备", "1份", "进水总管流量、化学需氧量、氨氮须自动监测并与监控平台联网"),
    ("21", "噪声排放情况说明与厂界噪声监测点位图", "必备", "1份", "执行 GB 12348-2008，按厂界功能区类别确定限值"),
    ("22", "无组织废气与有组织废气除臭设施说明", "必备", "1份", "含除臭装置排气筒高度、内径、排放速率限值依据 GB 14554-93"),
    ("23", "委托技术单位编制合同与资质证明", "按需", "1份", "受托技术机构不得弄虚作假；审批部门组织的技术评估不得向排污单位收费"),
    ("24", "排污许可证正本遗失、损毁情况说明", "按需", "1份", "补领时提交；已办理电子证照的可自行打印"),
]

CHECKLIST_REG = [
    ("1", "排污登记表（平台在线填报并生成回执）", "必备", "1份", "提交后即时生成登记编号与回执，无需提交纸质材料"),
    ("2", "营业执照或事业单位法人证书副本", "必备", "1份", "统一社会信用代码须与账号一致"),
    ("3", "实际运行维护主体说明", "必备", "1份", "无运营单位时由当前实际运行维护单位填报；责任主体变更后重新登记"),
    ("4", "主要产品信息（各站点设计规模与工艺）", "必备", "1份", "同一乡镇多个站点可在同一登记页面分别填报，备注站点名称与位置，获取一份回执"),
    ("5", "污染物排放去向与执行标准", "必备", "1份", "含排放口位置、排放规律、受纳水体"),
    ("6", "污染防治措施说明", "必备", "1份", "含工艺路线、消毒方式、污泥去向"),
]


def cmd_checklist(args):
    mgmt = args.management
    if mgmt == "registration":
        rows = CHECKLIST_REG
        title = "排污登记表填报资料清单"
    elif mgmt == "simple":
        rows = [r for r in CHECKLIST_COMMON if r[2] != "必备（重点管理）"]
        title = "简化管理排污许可证申报资料清单"
    else:
        rows = CHECKLIST_COMMON
        title = "重点管理排污许可证申报资料清单"

    if args.format == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(["序号", "资料名称", "是否必备", "份数", "备注"])
        for r in rows:
            w.writerow(list(r))
        _emit(buf.getvalue(), args.out)
        return 0

    if args.format == "xlsx":
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            sys.stderr.write(
                "[WARN] 未安装 openpyxl，xlsx 输出失败，回退为 CSV。安装命令：pip install openpyxl\n"
            )
            args.format = "csv"
            return cmd_checklist(args)
        wb = Workbook()
        ws = wb.active
        ws.title = "申报资料清单"
        header = ["序号", "资料名称", "是否必备", "份数", "备注"]
        ws.append(header)
        head_font = Font(bold=True, color="FFFFFF")
        head_fill = PatternFill("solid", fgColor="707070")
        for c in ws[1]:
            c.font = head_font
            c.fill = head_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
        for r in rows:
            ws.append(list(r))
        widths = [6, 46, 16, 8, 62]
        for i, w_ in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + i)].width = w_
        out = args.out or "permit_checklist.xlsx"
        if not out.lower().endswith(".xlsx"):
            out += ".xlsx"
        out_dir = os.path.dirname(os.path.abspath(out))
        if out_dir and not os.path.isdir(out_dir):
            os.makedirs(out_dir, exist_ok=True)
        wb.save(out)
        print("[OK] 已写入：%s（%d 行）" % (out, len(rows)))
        return 0

    # 默认 markdown
    lines = ["# %s" % title, ""]
    lines.append("> 依据：《排污许可管理条例》（国务院令第736号）第七条、第八条；"
                 "《排污许可管理办法》（生态环境部令第32号）第十八条；"
                 "HJ 978-2018 第4章。")
    lines.append("")
    lines.append("| 序号 | 资料名称 | 是否必备 | 份数 | 备注 |")
    lines.append("|---|---|---|---|---|")
    for r in rows:
        lines.append("| %s | %s | %s | %s | %s |" % r)
    lines.append("")
    lines.append("🔴 提交前逐项核对：注明「必备」的资料缺一即触发一次性补正告知，"
                 "受理时限自补正齐备之日起重新起算。")
    _emit("\n".join(lines), args.out)
    return 0


# ── 子命令 3：limit-calc ──────────────────────────────────────────
def cmd_limit_calc(args):
    pollutants = _split_list(args.pollutants)
    limits = [float(x) for x in _split_list(args.limits)]
    if not pollutants:
        pollutants = ["COD", "氨氮", "总氮", "总磷"]
    if not limits:
        limits = [_default_limit(p) for p in pollutants]
        if any(v <= 0 for v in limits):
            sys.stderr.write(
                "[WARN] 部分污染物无内置默认限值，已按 0 处理，须用 --limits 显式给出。\n"
            )
    if len(limits) < len(pollutants):
        sys.stderr.write("[ERROR] --limits 数量少于 --pollutants，按最小长度截断计算。\n")
        n = len(limits)
        pollutants = pollutants[:n]
    if len(limits) > len(pollutants):
        limits = limits[:len(pollutants)]

    env_batch = _split_list(args.env_batch)
    if env_batch and len(env_batch) < len(pollutants):
        sys.stderr.write("[ERROR] --env-batch 数量少于 --pollutants，缺失项不参与取严。\n")

    days = args.hours / 24.0
    flow = args.actual_flow if args.actual_flow else args.design_flow
    annual_volume = flow * days

    lines = []
    lines.append("# 许可排放量核算表（水污染物）")
    lines.append("")
    lines.append("## 一、核算参数")
    lines.append("")
    lines.append("| 参数 | 数值 | 口径 |")
    lines.append("|---|---|---|")
    lines.append("| 设计处理能力 | %.0f m³/d | 设计值 |" % args.design_flow)
    lines.append("| 水量基数取值 | %.0f m³/d | %s |" % (
        flow, "近三年实际排水量平均值" if args.actual_flow else "未投运/运行不满3年，取设计水量"))
    lines.append("| 年运行小时数 | %.0f h | 折合 %.2f 日 |" % (args.hours, days))
    lines.append("| 年排水量 Q | %.0f m³/a | Q = 水量基数 × 年运行日数 |" % annual_volume)
    lines.append("")
    lines.append("## 二、核算公式")
    lines.append("")
    lines.append("```")
    lines.append("E(j,许可) = Q × C(j,许可) × 10⁻⁶")
    lines.append("")
    lines.append("式中：")
    lines.append("  E(j,许可) —— 第 j 种水污染物年许可排放量，t/a")
    lines.append("  Q         —— 年排水量，m³/a（近三年实际排水量平均值；运行不满3年从投产之日起算；未投运取设计水量）")
    lines.append("  C(j,许可) —— 第 j 种水污染物许可排放浓度限值，mg/L")
    lines.append("  10⁻⁶      —— 单位换算系数（m³ × mg/L → t）")
    lines.append("")
    lines.append("取严规则：E(许可) = min( 按公式计算值, 环评批复文件明确的许可排放量 )")
    lines.append("          位于未达标重点区域、流域的，还须同时满足地方政府改善生态环境质量的特别要求")
    lines.append("```")
    lines.append("")
    lines.append("## 三、逐指标核算结果")
    lines.append("")
    lines.append("| 序号 | 污染物 | 许可浓度限值（mg/L） | 计算年许可量（t/a） | 环评批复量（t/a） | 取严后年许可量（t/a） | 取严依据 |")
    lines.append("|---|---|---|---|---|---|---|")
    total_calc = 0.0
    total_final = 0.0
    for i, (p, c) in enumerate(zip(pollutants, limits), 1):
        calc = annual_volume * c / 1e6
        env_v = None
        if env_batch and i - 1 < len(env_batch):
            env_v = float(env_batch[i - 1])
        if env_v is not None and env_v < calc:
            final = env_v
            basis_txt = "环评批复量从严"
        elif env_v is not None:
            final = calc
            basis_txt = "公式计算值从严"
        else:
            final = calc
            basis_txt = "无环评批复量，按公式计算值"
        total_calc += calc
        total_final += final
        lines.append("| %d | %s | %.4g | %.4f | %s | **%.4f** | %s |" % (
            i, _norm_pollutant(p), c, calc,
            ("%.4f" % env_v) if env_v is not None else "—",
            final, basis_txt,
        ))
    lines.append("| — | 合计 | — | %.4f | — | **%.4f** | — |" % (total_calc, total_final))
    lines.append("")
    lines.append("## 四、算例复算（口径校验）")
    lines.append("")
    lines.append("设计处理能力 5000 m³/d、年运行 8760 h、执行 GB 18918-2002 一级A：")
    lines.append("")
    lines.append("```")
    lines.append("Q = 5000 m³/d × 365 d = 1 825 000 m³/a")
    lines.append("E(化学需氧量) = 1 825 000 × 50 × 10⁻⁶ = 91.25 t/a")
    lines.append("E(氨氮)      = 1 825 000 × 8  × 10⁻⁶ = 14.60 t/a   # 限值 5(8)，低温期按 8 mg/L")
    lines.append("E(总氮)      = 1 825 000 × 15 × 10⁻⁶ = 27.375 t/a")
    lines.append("E(总磷)      = 1 825 000 × 0.5 × 10⁻⁶ = 0.9125 t/a → 0.913 t/a")
    lines.append("```")
    lines.append("")
    lines.append("🔴 核算结果须与环评批复文件、重点污染物总量控制指标逐项比对后取严；"
                 "比对未完成前不得将本表直接作为申请表附件提交。")
    _emit("\n".join(lines), args.out)
    return 0


# ── 子命令 4：form-template ───────────────────────────────────────
def cmd_form_template(args):
    ftype = args.type
    mgmt = args.management
    lines = []
    lines.append("# 排污许可证申请表填报模板")
    lines.append("")
    lines.append("> 模板字段以【待填：xxx】占位，逐项替换后提交。管理类别：%s。" % mgmt)
    lines.append("")
    lines.append("## 表1 排污单位基本信息表")
    lines.append("")
    lines.append("| 字段 | 填报内容 | 填报要求 |")
    lines.append("|---|---|---|")
    base_fields = [
        ("排污单位名称", "与营业执照完全一致；污水处理厂为独立法人的直接填厂名，非独立法人的填运营单位名称"),
        ("统一社会信用代码", "18位，须与平台实名认证信息一致"),
        ("注册地址", "与营业执照一致"),
        ("生产经营场所地址", "实际厂址；两个以上生产经营场所的分别申领"),
        ("行业类别", "污水处理及其再生利用（462）；水处理通用工序填112"),
        ("法定代表人（主要负责人）", "与营业执照一致"),
        ("技术负责人及联系电话", "负责持证日常管理的人员，须能联系到"),
        ("中心经纬度", "十进制度或度分秒，与监测点位图一致"),
        ("是否位于工业园区及园区名称", "工业园区配套污水处理设施须如实填报"),
        ("污水处理厂类型", "城镇污水处理厂 / 其他生活污水处理厂 / 工业废水集中处理厂"),
        ("排污许可证管理类别", mgmt),
        ("主要污染物类别", "废气 / 废水 / 工业固体废物 / 工业噪声，按实际勾选"),
        ("建设项目环评批复文号", "或环境影响登记表备案号；2015-01-01（含）后取得批复的须同时满足批复排放量要求"),
        ("排污权交易情况", "未发生填「未发生」；发生则填指标数量与交易凭证编号"),
    ]
    for name, req in base_fields:
        lines.append("| %s | 【待填：%s】 | %s |" % (name, name, req))
    lines.append("")
    lines.append("## 表2 主要产品及产能表")
    lines.append("")
    lines.append("| 生产线类别 | 名称及编号 | 设计处理能力 | 年运行小时数 | 厂外进水类别 | 备注 |")
    lines.append("|---|---|---|---|---|---|")
    lines.append("| 废水处理工程 | MF0001【待填：生产线名称】 | 【待填：设计处理能力 m³/d】 | 【待填：年运行小时数 h】 | "
                 "【待填：厂外生活污水/厂外工业废水/厂外雨水】 | 污水处理行业无产品产量指标，产能以设计处理能力体现 |")
    lines.append("| 固废处理工程 | MF0002【待填：污泥处理线名称】 | 【待填：处理能力 t/a】 | 【待填：年运行小时数 h】 | / | "
                 "【待填：是否外委处置】 |")
    lines.append("")
    lines.append("## 表3 原辅材料及燃料表")
    lines.append("")
    lines.append("| 工艺单元 | 药剂名称 | 浓度 | 设计投加量 | 投加点位编号 | 备注 |")
    lines.append("|---|---|---|---|---|---|")
    for u in ["混凝沉淀", "气浮", "深度处理（芬顿等高级氧化）", "消毒", "污泥脱水"]:
        lines.append("| %s | 【待填：药剂名称】 | 【待填：浓度百分比】 | 【待填：投加量 mg/L 或 kg/d】 | 【待填：TW/TA 编号】 | "
                     "【待填：与设计说明书一致】 |" % u)
    lines.append("")
    lines.append("## 表4 产排污环节与污染治理设施表")
    lines.append("")
    lines.append("| 类别 | 产污环节 | 编号 | 治理设施 | 编号 | 去向 |")
    lines.append("|---|---|---|---|---|---|")
    flows = [
        ("废水", "进水提升泵站", "TW001", "粗格栅", "TW004"),
        ("废水", "沉砂池", "TW002", "细格栅", "TW005"),
        ("废水", "调节池", "TW003", "提升泵", "TW006"),
        ("废水", "生化池（A²/O 或 A/O）", "TW008", "二沉池", "TW007"),
        ("废水", "深度处理（高效沉淀/微滤/芬顿）", "TW009", "消毒（次氯酸钠/紫外）", "TW010"),
        ("废水", "污泥脱水机房（压滤机）", "TS001", "事故池", "TW012"),
        ("废气", "格栅间、沉砂池、污泥浓缩脱水机房", "TA001", "生物除臭装置+排气筒", "TA006"),
        ("固废", "二沉池排泥", "TS001", "压滤脱水（含水率至 80% 及以下）", "—"),
    ]
    for row in flows:
        lines.append("| %s | 【待填：%s】 | %s | 【待填：%s】 | %s | 【待填：去向】 |" % (
            row[0], row[1], row[2], row[3], row[4]))
    lines.append("")
    lines.append("## 表5 排放口基本情况表")
    lines.append("")
    lines.append("| 编号 | 名称 | 类型 | 经纬度 | 排放去向 | 排放规律 | 许可内容 |")
    lines.append("|---|---|---|---|---|---|---|")
    lines.append("| DW001 | 废水总排口 | 主要排放口-总排口 | 【待填：经度/纬度】 | "
                 "【待填：直接排入xx河，入河排污口编号 RH001，批复文号】 | 连续排放 | 浓度限值 + 年许可排放量 |")
    lines.append("| DW002 | 雨水排放口 | 雨水排放口 | 【待填：经度/纬度】 | 【待填：汇入水体】 | 间断排放 | 仅浓度限值 |")
    lines.append("| DA001 | 除臭装置排气筒 | 一般排放口 | 【待填：经度/纬度】 | 大气环境 | 连续 | 仅浓度与速率限值，不许可排放量 |")
    lines.append("| 厂界 | 无组织（氨、硫化氢、臭气浓度、甲烷） | 无组织排放 | — | 大气环境 | — | 仅浓度限值 |")
    lines.append("| ZS001 | 厂界噪声（昼间/夜间） | 噪声 | 【待填：四至厂界点位】 | 环境 | 连续 | GB 12348-2008 对应功能区限值 |")
    lines.append("| MW001 | 进水总管（非排放口） | 进水监测点 | 【待填：经纬度】 | 厂内 | 连续 | 流量、化学需氧量、氨氮自动监测并联网 |")
    lines.append("")
    lines.append("## 表6 许可排放限值表")
    lines.append("")
    lines.append("| 污染物 | 许可排放浓度限值 | 执行标准 | 年许可排放量（t/a） | 确定依据 |")
    lines.append("|---|---|---|---|---|")
    for p, l in [("化学需氧量（CODcr）", "50 mg/L"), ("氨氮（以N计）", "5(8) mg/L"),
                 ("总氮（以N计）", "15 mg/L"), ("总磷（以P计）", "0.5 mg/L")]:
        lines.append("| %s | 【待填：%s】 | GB 18918-2002 一级A | 【待填：t/a】 | "
                     "浓度取标准限值；排放量按 Q×C×10⁻⁶ 与环评批复量取严 |" % (p, l))
    lines.append("| pH值 | 6～9 | GB 18918-2002 | 不许可排放量 | 仅浓度 |")
    lines.append("| 悬浮物 | 10 mg/L | GB 18918-2002 | 不许可排放量 | 仅浓度 |")
    lines.append("| 五日生化需氧量 | 10 mg/L | GB 18918-2002 | 不许可排放量 | 仅浓度 |")
    lines.append("| 氨（有组织） | 【待填：浓度限值】 | GB 14554-93 | 不许可排放量 | 一般排放口仅浓度与速率 |")
    lines.append("| 硫化氢（有组织） | 【待填：浓度限值】 | GB 14554-93 | 不许可排放量 | 一般排放口仅浓度与速率 |")
    lines.append("| 氨（厂界无组织） | 1.5 mg/m³ | GB 14554-93 | 不许可排放量 | 厂界无组织仅浓度 |")
    lines.append("| 硫化氢（厂界无组织） | 0.06 mg/m³ | GB 14554-93 | 不许可排放量 | 厂界无组织仅浓度 |")
    lines.append("| 臭气浓度（厂界无组织） | 20（无量纲） | GB 14554-93 | 不许可排放量 | 厂界无组织仅浓度 |")
    lines.append("| 甲烷（厂区体积浓度最高处） | 1% | GB 18918-2002 | 不许可排放量 | 厂内无组织 |")
    lines.append("")
    lines.append("## 表7 自行监测方案表")
    lines.append("")
    lines.append("| 监测点位 | 监测指标 | 监测方式 | 监测频次 | 分析方法标准 |")
    lines.append("|---|---|---|---|---|")
    lines.append("| 进水总管 MW001 | 流量、化学需氧量、氨氮 | 自动 | 连续 | 自动监测仪，与监控平台联网 |")
    lines.append("| 进水总管 MW001 | 总磷、总氮 | 手工 | 1次/日 | 【待填：方法标准号】 |")
    lines.append("| 废水总排口 DW001 | 流量、pH值、水温、化学需氧量、氨氮、总磷、总氮 | 自动 | 连续 | 自动监测仪，与监控平台联网 |")
    lines.append("| 废水总排口 DW001 | 悬浮物 | 手工，瞬时采样不少于3个 | 1次/日 | GB 11901-1989 |")
    lines.append("| 废水总排口 DW001 | 五日生化需氧量 | 手工 | 1次/月 | HJ 505-2009 |")
    lines.append("| 废水总排口 DW001 | 色度、石油类、动植物油、阴离子表面活性剂 | 手工 | 1次/月 | 【待填：方法标准号】 |")
    lines.append("| 废水总排口 DW001 | 总镉、总铬、总汞、总铅、总砷、六价铬 | 手工，混合采样 | 1次/月 | 【待填：方法标准号】 |")
    lines.append("| 废水总排口 DW001 | 粪大肠菌群 | 手工 | 1次/季 | HJ 755-2015 |")
    lines.append("| 雨水排放口 DW002 | pH值、化学需氧量、氨氮、悬浮物 | 手工，瞬时采样不少于3个 | 1次/日（监测一年无异常可放宽至每季度1次） | 【待填：方法标准号】 |")
    lines.append("| 除臭排气筒 DA001 | 臭气浓度、氨、硫化氢 | 手工 | 1次/半年 | HJ 1262-2022、HJ 533-2009、GB/T 14678-1993 |")
    lines.append("| 厂界 | 氨、硫化氢、臭气浓度 | 手工 | 1次/半年 | GB/T 14675-1993 等 |")
    lines.append("| 厂界 | 环境噪声（昼间、夜间） | 手工 | 1次/季 | GB 12348-2008 |")
    lines.append("")
    lines.append("## 表8 环境管理台账要求表")
    lines.append("")
    lines.append("| 台账类别 | 记录内容 | 记录频次 | 保存形式与期限 |")
    lines.append("|---|---|---|---|")
    lines.append("| 基本信息 | 生产设施与污染防治设施技术参数及设计值；防渗漏措施落实情况 | 无变化 1次/年，有变化时及时记录 | 电子+纸质，不少于5年 |")
    lines.append("| 生产设施运行 | 正常工况运行状态、生产负荷、进出水量；非正常工况起止时间、原因、措施 | 1次/日或批次；异常按工况期 | 电子+纸质，不少于5年 |")
    lines.append("| 污染防治设施运行 | 运行情况、处理效率、药剂添加量、污泥产生量及含水率、异常情况 | 1次/日；异常按异常期 | 电子+纸质，不少于5年 |")
    lines.append("| 监测记录 | 手工监测记录、自动监测运维记录、校准校验记录、同步生产工况 | 按 HJ 819-2017 与 HJ 1106-2020 | 电子+纸质，不少于5年 |")
    lines.append("| 其他环境管理信息 | 无组织废气防治措施维护、特殊时段管理要求执行情况、固废与危废去向 | 按法规与许可证规定 | 电子+纸质，不少于5年 |")
    lines.append("")
    lines.append("## 表9 执行报告要求表")
    lines.append("")
    lines.append("| 报告类型 | 报告周期 | 主要内容 | 上报要求 |")
    lines.append("|---|---|---|---|")
    lines.append("| 年度执行报告 | 持证超过3个月的年度为当年全年；不足3个月纳入下一年度 | "
                 "基本情况、污染防治设施运行、自行监测执行、台账执行、实际排放及合规判定、信息公开、内部环境管理体系、结论、附图附件 | "
                 "按许可证规定时间在平台提交；电子版与签字盖章书面件一致 |")
    lines.append("| 季度执行报告 | 持证超过1个月的季度为当季全季；不足1个月纳入下一季度 | "
                 "污染物实际排放浓度与排放量、合规判定分析、超标或设施异常说明；"
                 "并含各月生产小时数、处理水量、药剂消耗量、新水用量与废水排放量 | "
                 "按许可证规定时间在平台提交 |")
    lines.append("| 月度执行报告 | 持证超过10日的月份为当月全月 | 实际排放浓度与排放量、异常说明 | 地方要求时提交 |")
    lines.append("")
    if ftype == "registration":
        lines.append("## 表10 排污登记表填报要点")
        lines.append("")
        lines.append("| 填报项 | 填报内容 | 要求 |")
        lines.append("|---|---|---|")
        lines.append("| 单位基本信息 | 【待填：单位名称、统一社会信用代码、生产经营场所地址】 | 与实际运行维护主体一致 |")
        lines.append("| 主要产品信息 | 【待填：各站点设计规模与工艺】 | 同一乡镇多个站点可分列，备注站点名称与位置 |")
        lines.append("| 污染物排放去向 | 【待填：排放口位置、排放规律、受纳水体】 | 与现场一致 |")
        lines.append("| 执行排放标准 | 【待填：标准号与限值】 | GB 18918-2002 或地方标准 |")
        lines.append("| 采取的污染防治措施 | 【待填：工艺路线与消毒方式】 | 与设计说明书一致 |")
        lines.append("")
        lines.append("🔴 登记表提交后即时生成登记编号与回执；信息变动自变动之日起20日内变更登记；"
                     "依法须申领排污许可证的，须及时申领并注销登记表。")
    else:
        lines.append("## 表10 附图附件清单")
        lines.append("")
        lines.append("| 序号 | 附件名称 | 形式要求 |")
        lines.append("|---|---|---|")
        atts = [
            "厂区总平面布置图（标注产排污环节与排放口编号）",
            "污水处理工艺流程图（标注 TW/TA/TS 编号）",
            "监测点位示意图（含进水总管、废水总排口、雨水排放口、除臭排气筒、厂界噪声点位）",
            "排放口规范化建设及标志牌照片",
            "管网走向与纳污范围示意图",
            "环评批复文件或环境影响登记表备案回执",
            "统一社会信用代码证与法定代表人身份证明",
            "许可排放量限值计算过程",
            "总量控制指标来源与削减替代说明",
            "近三年实际排水量与运行小时数统计表",
        ]
        for i, a in enumerate(atts, 1):
            lines.append("| %d | 【待填：%s】 | PDF，单个附件不超过平台限制大小，须清晰可辨 |" % (i, a))
    lines.append("")
    lines.append("🔴 模板中全部【待填：xxx】占位须替换完毕方可提交；"
                 "提交前须由技术负责人逐表复核一次并留存复核记录。")
    _emit("\n".join(lines), args.out)
    return 0


# ── 参数解析 ──────────────────────────────────────────────────────
def build_parser():
    p = argparse.ArgumentParser(
        prog="permit_form_builder.py",
        description="排污许可证申领及填报助手：管理类别判定、申报资料清单、许可排放量核算、申请表填报模板。",
        epilog="依据：国务院令第736号、生态环境部令第32号、名录（2019年版，部令第11号）、HJ 978-2018、HJ 944-2018。",
    )
    p.add_argument("--version", action="version", version="permit_form_builder.py %s" % VERSION)
    sub = p.add_subparsers(dest="command", metavar="{classify,checklist,limit-calc,form-template}")

    # classify
    c = sub.add_parser("classify", help="管理类别判定（重点/简化/登记）")
    c.add_argument("--industry", default="污水处理及其再生利用",
                   help="行业类别，默认「污水处理及其再生利用」；可填 462 / 112 / 水处理通用工序")
    c.add_argument("--scale", type=float, default=0.0, help="设计处理能力，单位 m³/d（城乡污水集中处理场所按此分档）")
    c.add_argument("--water-type", dest="water_type", default="domestic",
                   choices=["domestic", "industrial", "mixed"],
                   help="污水来源：domestic=城乡生活污水集中处理场所，industrial=工业废水集中处理场所，mixed=混合")
    c.add_argument("--discharge", default="direct", choices=["direct", "indirect", "reuse"],
                   help="排放去向：direct=直接排入环境水体，indirect=间接排放，reuse=再生利用不外排")
    c.add_argument("--cod-total", dest="cod_total", type=float, default=None,
                   help="化学需氧量年排放量 t/a（用于核对名录第七条是否适用）")
    c.add_argument("--out", default=None, help="输出 Markdown 文件路径")
    c.set_defaults(func=cmd_classify)

    # checklist
    ck = sub.add_parser("checklist", help="生成申报资料清单")
    ck.add_argument("--management", default="key", choices=["key", "simple", "registration"],
                    help="管理类别：key=重点管理，simple=简化管理，registration=登记管理")
    ck.add_argument("--format", default="md", choices=["md", "csv", "xlsx"],
                    help="输出格式，默认 md；xlsx 需 openpyxl，缺失时自动降级为 csv")
    ck.add_argument("--out", default=None, help="输出文件路径")
    ck.set_defaults(func=cmd_checklist)

    # limit-calc
    lc = sub.add_parser("limit-calc", help="许可排放量核算")
    lc.add_argument("--design-flow", dest="design_flow", type=float, required=True,
                    help="设计处理能力，单位 m³/d")
    lc.add_argument("--hours", type=float, default=8760.0, help="年运行小时数，默认 8760")
    lc.add_argument("--pollutants", default="COD,氨氮,总氮,总磷",
                    help="污染物，逗号分隔，默认 COD,氨氮,总氮,总磷")
    lc.add_argument("--limits", default=None,
                    help="对应的许可排放浓度限值 mg/L，逗号分隔；缺省时按 GB 18918-2002 一级A 取值")
    lc.add_argument("--actual-flow", dest="actual_flow", type=float, default=None,
                    help="近三年实际排水量平均值 m³/d；给定后替代设计水量作为水量基数")
    lc.add_argument("--env-batch", dest="env_batch", default=None,
                    help="环评批复文件明确的许可排放量 t/a，逗号分隔，与 --pollutants 一一对应；用于取严")
    lc.add_argument("--out", default=None, help="输出 Markdown 文件路径")
    lc.set_defaults(func=cmd_limit_calc)

    # form-template
    ft = sub.add_parser("form-template", help="生成申请表填报模板")
    ft.add_argument("--type", dest="type", default="wastewater",
                    choices=["wastewater", "industrial", "registration"],
                    help="模板类型：wastewater=城乡污水处理厂，industrial=工业废水集中处理厂，registration=排污登记表")
    ft.add_argument("--management", default="key", choices=["key", "simple"],
                    help="管理类别：key=重点管理，simple=简化管理")
    ft.add_argument("--out", default=None, help="输出 Markdown 文件路径")
    ft.set_defaults(func=cmd_form_template)

    return p


def main(argv=None):
    parser = build_parser()
    args = parser.parse_args(argv)
    if not getattr(args, "command", None):
        parser.print_help()
        return 0
    try:
        return args.func(args)
    except Exception as exc:  # noqa
        sys.stderr.write("[ERROR] %s: %s\n" % (type(exc).__name__, exc))
        return 1


if __name__ == "__main__":
    sys.exit(main())
